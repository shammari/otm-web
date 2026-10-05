"""Result exports: the files the legacy Save buttons wrote, from Python results.

==========================================  =========================================================
Legacy writer (MATLAB)                      Here
==========================================  =========================================================
``OutputTextFile``                          :func:`format_indices_text` -> ``<name>_Output_Text_File.txt``
``FiberOutputExcelFile`` (.xls)             :func:`write_table_xlsx` -> ``<name>_FiberIndicesOutput.xlsx``
``saveIndexFigure``                         :func:`index_distribution_figure` -> ``<name>_<Index>_Distribution.png``
``OutputPO2TextFile``                       :func:`format_po2_text` -> ``<name>_PO2_Stats.txt``
``OutputCardiacPO2TextFile``                same, cardiac layout -> ``<name>_PO2_Stats_Output_Text_File.txt``
``FiberPO2OutputExcelsheet`` (.xls)         ``<name>_FiberPO2Stats.xlsx``
``FEMSimulation`` Save (pdeplot)            :func:`profile_figure` -> ``<name>_PO2_Profile.png``
``PlotOxygenDistribution``                  :func:`density_figure` -> ``<name>_PO2_ProbabilityDensity.png``
``PlotOxygenUptakeDistribution``            ``<name>_VO2_ProbabilityDensity.png``, ``<name>_VO2_Profile.png``
``OxygenFluxGUI`` Save                      :func:`flux_figure` -> ``<name>_PO2_Flux_Lines.png``
(new)                                       :func:`write_flux_csv` -> ``<name>_PO2_Flux_Lines.csv``
==========================================  =========================================================

The text files reproduce the legacy ``fprintf`` layouts byte for byte (CRLF line
ends, the same column widths, ``NaN``/``Inf`` spelt as MATLAB does), so scripts
that parse the old files keep working. The indices file uses the August-2022
layout (with the ``Dmax`` column). Spreadsheets are ``.xlsx`` (the ``.xls``
format of ``writetable`` is no longer written by current Excel/openpyxl).

Figures use matplotlib's object API with the Agg canvas (no pyplot state), so
they can be drawn on a worker thread.

Everything here is pure: no Qt. The high-level ``export_*`` functions take the
core result objects and a folder and return the list of files written.
"""

from __future__ import annotations

import csv
import math
import os
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np

from .progress import ProgressCallback, report

__all__ = [
    "IndicesReport", "FIBER_INDEX_COLUMNS", "FIBER_PO2_COLUMNS", "INDEX_FIGURES",
    "mformat", "format_indices_text", "format_po2_text", "write_text", "write_table_xlsx", "read_table_xlsx",
    "histc", "matlab_smooth", "density_curves", "fiber_po2_table",
    "index_distribution_figure", "density_figure", "profile_figure", "flux_figure", "write_flux_csv",
    "export_indices", "export_po2", "export_flux", "safe_base_name",
]

CRLF = "\r\n"


# ==========================================================================
# MATLAB-compatible number formatting
# ==========================================================================


def mformat(fmt: str, value: Any) -> str:
    """One ``fprintf`` conversion (``%-8.2f``, ``%-8d``, ``%1.4e``, ``%-13s`` ...).

    Non-finite numbers are written ``NaN``/``Inf``/``-Inf`` with the field width
    kept, as MATLAB does; ``%d`` of an integral float prints the integer.
    """
    conv = fmt[-1]
    if conv == "s":
        return fmt % str(value)
    v = float(value)
    if not math.isfinite(v):
        txt = "NaN" if math.isnan(v) else ("Inf" if v > 0 else "-Inf")
        spec = fmt[1:-1].split(".")[0]                # flags + width
        width = int(spec.lstrip("-+ #0") or 0)
        return txt.ljust(width) if "-" in spec else txt.rjust(width)
    if conv == "d":
        if v != int(v):                               # MATLAB switches to %e for non-integers
            return mformat(fmt[:-1] + "e", v)
        return fmt % int(v)
    return fmt % v


def _row(fmts: Sequence[str], values: Sequence[Any], sep: str = " ") -> str:
    return sep.join(mformat(f, v) for f, v in zip(fmts, values))


def _num2str(v: float) -> str:
    """``num2str(v, '%1.4e')``."""
    return mformat("%1.4e", v)


# ==========================================================================
# indices
# ==========================================================================


@dataclass
class IndicesReport:
    """Everything ``OutputTextFile`` / ``FiberOutputExcelFile`` / ``saveIndexFigure`` use.

    Lists are per ROI capillary (``domain_areas`` ... ``sf``) and per ROI fibre
    (``fiber_no`` ... ``cfpe``); ``fiber_no`` is 1-based into the full fibre list
    and ``fiber_type`` holds the raw codes (1, 21, 22, 0). Fibre fields are empty
    for cardiac tissue (no supply indices), which selects the cardiac layout.
    """

    n_capillaries: int
    capillary_density: float            # per mm^2 (legacy label says um^{-2})
    x_length_um: float
    y_length_um: float
    n_domains: int
    n_domains_in_roi: int
    area_mean: float
    area_std: float
    area_sem: float
    diameter_mean: float
    diameter_std: float
    diameter_sem: float
    nn_mean_of_means: float
    nn_std_of_means: float
    nn_sem_of_means: float
    nn_mean_of_stds: float
    nn_mean_of_sems: float
    domain_areas: np.ndarray
    domain_diameters: np.ndarray
    nn_min: np.ndarray
    nn_unique_min: np.ndarray
    nn_mean: np.ndarray
    nn_random: np.ndarray
    nn_all: np.ndarray
    fdr: np.ndarray = field(default_factory=lambda: np.zeros(0))
    sf: np.ndarray = field(default_factory=lambda: np.zeros(0))
    fiber_no: np.ndarray = field(default_factory=lambda: np.zeros(0, int))
    fiber_type: np.ndarray = field(default_factory=lambda: np.zeros(0, int))
    fiber_area: np.ndarray = field(default_factory=lambda: np.zeros(0))
    dfr: np.ndarray = field(default_factory=lambda: np.zeros(0))
    lcfr: np.ndarray = field(default_factory=lambda: np.zeros(0))
    lcd: np.ndarray = field(default_factory=lambda: np.zeros(0))
    dmax: np.ndarray = field(default_factory=lambda: np.zeros(0))
    cc: np.ndarray = field(default_factory=lambda: np.zeros(0))
    cfi: np.ndarray = field(default_factory=lambda: np.zeros(0))
    perimeter: np.ndarray = field(default_factory=lambda: np.zeros(0))
    cfpe: np.ndarray = field(default_factory=lambda: np.zeros(0))
    has_fibres: bool = True

    @property
    def log_sd_area(self) -> float:
        return _std1(np.log10(self.domain_areas))

    @property
    def log_sd_diameter(self) -> float:
        return _std1(np.log10(self.domain_diameters))

    # ---- constructors ----------------------------------------------------------------------------
    @classmethod
    def from_morphometrics(cls, m: Any, fiber_types: Optional[Sequence[int]] = None) -> "IndicesReport":
        """From :class:`otm_core.models.Morphometrics`; ``fiber_types`` = raw codes of
        *all* geometry fibres (``Geometry.fiber_types``)."""
        d, nn, s = m.domains, m.nearest_neighbours, m.supply
        S = nn.summary
        kw: Dict[str, Any] = dict(
            n_capillaries=int(d.total_num_capillaries), capillary_density=float(d.capillary_density_per_mm2),
            x_length_um=float(d.x_length_um), y_length_um=float(d.y_length_um),
            n_domains=int(len(d.capillary_domain_areas_um2)), n_domains_in_roi=int(d.n_domains_in_roi),
            area_mean=d.area_stats.mean, area_std=d.area_stats.std, area_sem=d.area_stats.sem,
            diameter_mean=d.diameter_stats.mean, diameter_std=d.diameter_stats.std,
            diameter_sem=d.diameter_stats.sem,
            nn_mean_of_means=S["MeanOfMeans"], nn_std_of_means=S["STDOfMeans"], nn_sem_of_means=S["SEMOfMeans"],
            nn_mean_of_stds=S["MeanOfSTDs"], nn_mean_of_sems=S["MeanOfSEMs"],
            domain_areas=_a(d.roi_capillary_domain_areas_um2), domain_diameters=_a(d.domain_equivalent_diameter_um),
            nn_min=_a(nn.mins), nn_unique_min=_a(nn.unique_min_list), nn_mean=_a(nn.means),
            nn_random=_a(nn.rnd_list), nn_all=_a(nn.all_distances))
        if s is None:
            return cls(**kw, has_fibres=False)
        idx = np.asarray(s.roi_fiber_index, int)
        types = np.asarray(fiber_types, int)[idx] if fiber_types is not None and len(fiber_types) \
            else np.zeros(idx.size, int)
        return cls(**kw, fdr=_a(s.FDR.values), sf=_a(s.SF.values), fiber_no=idx + 1, fiber_type=types,
                   fiber_area=_a(s.fiber_area_um2), dfr=_a(s.DFR.values), lcfr=_a(s.LCFR.values),
                   lcd=_a(s.LCD.values), dmax=_a(s.Dmax.values), cc=_a(s.CC.values), cfi=_a(s.CFi.values),
                   perimeter=_a(s.FPi.values), cfpe=_a(s.CFPE.values), has_fibres=True)

    @classmethod
    def from_legacy(cls, raw: Mapping[str, Any], fiber_types: Optional[Sequence[int]] = None) -> "IndicesReport":
        """From a MATLAB ``Morphometrics`` struct (as in ``baseline_outputs.json``,
        ``null`` = NaN); used to check the writers against MATLAB's own numbers."""
        cd, nn = raw["CapDomains"], raw["NearestNeighbors"]
        kw: Dict[str, Any] = dict(
            n_capillaries=int(cd["TotalNumOfCapillaries"]), capillary_density=float(cd["CapillaryDensity"]),
            x_length_um=float(cd["XLengthScale"]), y_length_um=float(cd["YLengthScale"]),
            n_domains=int(cd["TotalNumOfCapDomains"]), n_domains_in_roi=int(cd["NumberCapillaryDomainsInROI"]),
            area_mean=_f(cd["MeanArea"]), area_std=_f(cd["STDArea"]), area_sem=_f(cd["SEMArea"]),
            diameter_mean=_f(cd["MeanDomainEquivalentDiameter"]), diameter_std=_f(cd["STDDomainEquivalentDiameter"]),
            diameter_sem=_f(cd["SEMDomainEquivalentDiameter"]),
            nn_mean_of_means=_f(nn["MeanOfMeans"]), nn_std_of_means=_f(nn["STDOfMeans"]),
            nn_sem_of_means=_f(nn["SEMOfMeans"]), nn_mean_of_stds=_f(nn["MeanOfSTDs"]),
            nn_mean_of_sems=_f(nn["MeanOfSEMs"]),
            domain_areas=_a(cd["ROICapillaryDomainAreas"]), domain_diameters=_a(cd["DomainEquivalentDiameter"]),
            nn_min=_a(nn["MinList"]), nn_unique_min=_a(nn["UniqueMinList"]), nn_mean=_a(nn["MeansList"]),
            nn_random=_a(nn["RndList"]), nn_all=_a(nn["AllDistances"]))
        si = raw.get("SupplyIndices")
        if not si:
            return cls(**kw, has_fibres=False)
        L = lambda k: _a(si[k]["LIST"])  # noqa: E731
        no = _a(si["Fiber"]["List"]).astype(int)
        types = np.asarray(fiber_types, int)[no - 1] if fiber_types is not None else np.zeros(no.size, int)
        return cls(**kw, fdr=L("FDR"), sf=L("SF"), fiber_no=no, fiber_type=types, fiber_area=_a(si["Fiber"]["Area"]),
                   dfr=L("DFR"), lcfr=L("LCFR"), lcd=L("LCD"), dmax=L("Dmax"), cc=L("CC"), cfi=L("CFi"),
                   perimeter=L("FPi"), cfpe=L("CFPE"), has_fibres=True)

    # ---- tables ------------------------------------------------------------------------------------
    def fiber_table(self) -> Dict[str, np.ndarray]:
        """``FiberOutputExcelFile`` columns."""
        return {"Fiber_No": self.fiber_no, "Fiber_Type": self.fiber_type, "Fiber_Area_um2": self.fiber_area,
                "Perimeter_um": self.perimeter, "DFR": self.dfr, "LCFR": self.lcfr, "LCD": self.lcd,
                "Dmax": self.dmax, "CCi": self.cc, "CFi": self.cfi, "CFPE": self.cfpe}

    def distributions(self) -> List[Tuple[str, str, np.ndarray]]:
        """(name, units, values) of every ``saveIndexFigure`` panel, legacy order."""
        out = []
        for name, units, attr in INDEX_FIGURES:
            v = getattr(self, attr)
            if attr in _FIBRE_FIELDS and not self.has_fibres:
                continue
            out.append((name, units, np.asarray(v, float)))
        return out


FIBER_INDEX_COLUMNS = ("Fiber_No", "Fiber_Type", "Fiber_Area_um2", "Perimeter_um", "DFR", "LCFR", "LCD", "Dmax",
                       "CCi", "CFi", "CFPE")
FIBER_PO2_COLUMNS = ("Fiber_No", "Fiber_Type", "Average_PO2_mmHg", "StDev_PO2_mmHg", "Hypoxic_Fraction")

INDEX_FIGURES: Tuple[Tuple[str, str, str], ...] = (
    ("Fiber area", " (um sq)", "fiber_area"), ("DFR", "", "dfr"), ("FDR", "", "fdr"), ("LCFR", "", "lcfr"),
    ("LCD", " (1/mm sq)", "lcd"), ("Equivalent diameter", " (um)", "domain_diameters"),
    ("Capillary domain Area", " (um sq)", "domain_areas"), ("Mean NN distance", " (um)", "nn_mean"),
    ("Minimum NN distance", " (um)", "nn_min"), ("Unique minimum NN distance", " (um)", "nn_unique_min"),
    ("Unique random NN distance", " (um)", "nn_random"), ("All NN distances", " (um)", "nn_all"),
)
"""(name, units, IndicesReport field) as in ``IndicesMenu``. The legacy LCD axis
said ``(1/um sq)``; the values are per mm^2, so the label is corrected here."""
_FIBRE_FIELDS = {"fiber_area", "dfr", "fdr", "lcfr", "lcd"}


def _a(v: Any) -> np.ndarray:
    if v is None:
        return np.zeros(0)
    arr = np.asarray([np.nan if x is None else x for x in np.ravel(np.asarray(v, dtype=object))], dtype=float) \
        if isinstance(v, (list, tuple)) else np.asarray(v, dtype=float).ravel()
    return arr


def _f(v: Any) -> float:
    return float("nan") if v is None else float(v)


def _std1(v: np.ndarray) -> float:
    v = np.asarray(v, float)
    return float(np.std(v, ddof=1)) if v.size > 1 else (0.0 if v.size == 1 else float("nan"))


def format_indices_text(rep: IndicesReport) -> str:
    """``OutputTextFile`` (August-2022 layout, with ``Dmax``); CRLF line ends."""
    line = "%-30s"
    out: List[str] = []

    def block(rows: Sequence[Tuple[str, Any, str]], int_rows: Iterable[int]) -> None:
        ints = set(int_rows)
        for k, (label, value, unit) in enumerate(rows):
            num = "%-8d" if k in ints else "%-8.2f"
            out.append(_row((line, num, "%2s"), (label, value, unit)) + CRLF + (CRLF if k == len(rows) - 1 else ""))

    block([("Number of Capillaries", rep.n_capillaries, ""),
           ("Capillary Density", rep.capillary_density, "um^{-2}"),
           ("X Length Scale", rep.x_length_um, "um"),
           ("Y Length Scale", rep.y_length_um, "um")], (0,))
    block([("Number of Voronoi Cells", rep.n_domains, ""),
           ("Included Voronoi Cells", rep.n_domains_in_roi, ""),
           ("Mean[Voronoi Cell Area]", rep.area_mean, "um^2"),
           ("STD[Voronoi Cell Area]", rep.area_std, "um^2"),
           ("SEM[Voronoi Cell Area]", rep.area_sem, "um^2"),
           ("LogSD[Voronoi Cell Area]", rep.log_sd_area, "")], (0, 1))
    block([("Mean[Domain Eq. Diameter]", rep.diameter_mean, "um"),
           ("STD[Domain Eq. Diameter]", rep.diameter_std, "um"),
           ("SEM[Domain Eq. Diameter]", rep.diameter_sem, "um"),
           ("LogSD[Domain Eq. Diameter]", rep.log_sd_diameter, ""),
           ("Mean[Mean NN Distance]", rep.nn_mean_of_means, "um"),
           ("STD[Mean NN Distance]", rep.nn_std_of_means, "um"),
           ("SEM[Mean NN Distance]", rep.nn_sem_of_means, "um"),
           ("Mean[STD NN Distance]", rep.nn_mean_of_stds, "um"),
           ("Mean[SEM NN Distance]", rep.nn_mean_of_sems, "um")], ())

    cap_head = ("Domain Areas (um^2)", "Min NN (um)", "Unique Min NN (um)", "Mean NN (um)", "Rand NN (um)")
    cap_cols = [rep.domain_areas, rep.nn_min, rep.nn_unique_min, rep.nn_mean, rep.nn_random]
    if not rep.has_fibres:                                                   # cardiac (nargin == 2)
        out.append(_row(("%-22s", "%-14s", "%-21s", "%-15s", "%-12s"), cap_head) + CRLF + CRLF)
        fm = ("%-22.2f", "%-14.2f", "%-21.2f", "%-15.2f", "%-12.2f")
        out.extend(_row(fm, r) + CRLF for r in _rows(cap_cols))
    else:
        out.append(_row(("%-22s", "%-14s", "%-21s", "%-15s", "%-15s", "%-8s", "%-8s"), cap_head + ("FDR", "SF"))
                   + CRLF + CRLF)
        fm = ("%-22.2f", "%-14.2f", "%-21.2f", "%-15.2f", "%-15.2f", "%-8.0f", "%-8.0f")
        out.extend(_row(fm, r) + CRLF for r in _rows(cap_cols + [rep.fdr, rep.sf]))
        head = ("Fiber No.", "Fiber Type", "Fiber Area (um^2)", "DFR", "LCFR", "LCD", "Dmax", "CC", "CFi", "Pi", "CFPE")
        out.append(CRLF + _row(("%-12s", "%-13s", "%-20s", "%-8s", "%-8s", "%-11s", "%-10s", "%-8s", "%-8s", "%-8s",
                                "%-8s"), head) + CRLF + CRLF)
        fm = ("%-12.0f", "%-13.0f", "%-20.2f", "%-8.0f", "%-8.2f", "%-11.2f", "%-10.2f", "%-8.0f", "%-8.2f", "%-8.2f",
              "%-8.2f")
        out.extend(_row(fm, r) + CRLF for r in _rows([rep.fiber_no, rep.fiber_type, rep.fiber_area, rep.dfr, rep.lcfr,
                                                       rep.lcd, rep.dmax, rep.cc, rep.cfi, rep.perimeter, rep.cfpe]))
    out.append(CRLF + mformat("%-15s", "All NN (um)") + CRLF + CRLF)
    out.extend(mformat("%-15.2f", v) + CRLF for v in rep.nn_all)
    return "".join(out)


def _rows(cols: Sequence[np.ndarray]) -> List[Tuple[float, ...]]:
    """``fprintf(fmt, [c1, c2, ...]')``: one line per row. MATLAB would refuse
    columns of different length (horzcat); here they must agree too."""
    n = {len(c) for c in cols}
    if len(n) > 1:
        raise ValueError(f"columns of different length: {[len(c) for c in cols]}")
    return list(zip(*cols))


# ==========================================================================
# PO2 / MO2 statistics
# ==========================================================================

_PO2_LABELS = (("Interstitia", "Interstitia"), ("FiberI", "Fiber I    "), ("FiberIIa", "Fiber IIa  "),
               ("FiberIIb", "Fiber IIb  "), None, ("AllFibers", "All Fibers "), ("Tissue", "Tissue     "))
_MO2_LABELS = (("Interstitia", "Interstitia"), ("FiberI", "Fiber I"), ("FiberIIa", "Fiber IIa"),
               ("FiberIIb", "Fiber IIb"), None, ("AllFibers", "All Fibers"), ("Tissue", "Tissue"))


def format_po2_text(po2: Mapping[str, Mapping[str, float]], mo2: Mapping[str, Mapping[str, float]],
                    fibers: Optional[Mapping[str, np.ndarray]] = None, tissue: str = "skeletal") -> str:
    """``OutputPO2TextFile`` (skeletal, with the per-fibre table when ``fibers``
    is given) or ``OutputCardiacPO2TextFile`` (tissue rows only).

    ``po2``/``mo2``: :func:`otm_core.fem.po2_statistics` / ``mo2_statistics``.
    ``fibers``: :func:`fiber_po2_table`.
    """
    cardiac = tissue == "cardiac"
    pf = ("%-13s", "%-16s", "%-17s", "%-10s")
    mf = ("%-13s", "%-17s", "%-12s", "%-12s", "%-19s")
    out = ["PO2 predictions " + CRLF + CRLF]
    rows: List[Tuple[str, ...]] = [("Compartment", "Mean PO2 (mmHg)", "StDev PO2 (mmHg)", "% Hypoxia"),
                                   ("-----------", "---------------", "----------------", "-----------")]
    for item in ((("Tissue", "Tissue     "),) if cardiac else _PO2_LABELS):
        if item is None:
            rows.append(rows[1])
            continue
        v = po2.get(item[0], {})
        rows.append((item[1], _num2str(v.get("mean_mmHg", np.nan)) + "    ", _num2str(v.get("std_mmHg", np.nan)),
                     _num2str(v.get("hypoxia_pct", np.nan))))
    out.extend(_row(pf, r) + " " + CRLF for r in rows)
    out.append(CRLF + CRLF + "MO2 predictions " + CRLF + CRLF)
    rows = [("Compartment", "VO2max (ml/ml*s)", "Mean MO2", "StDev MO2", "% MO2 < 0.5 VO2max"),
            ("-----------", "----------------", "-----------", "-----------", "------------------")]
    for item in ((("Tissue", "Tissue     "),) if cardiac else _MO2_LABELS):
        if item is None:
            rows.append(("-----------", "-------------", "-----------", "-----------", "------------------"))
            continue
        v = mo2.get(item[0], {})
        rows.append((item[1], _num2str(v.get("vo2max", np.nan)) + "  ", _num2str(v.get("mean_mo2", np.nan)),
                     _num2str(v.get("std_mo2", np.nan)), _num2str(v.get("pct_below_half_vo2max", np.nan))))
    end = " " + CRLF if cardiac else CRLF
    out.extend(_row(mf, r) + end for r in rows)
    if not cardiac and fibers is not None:
        out.append(CRLF + CRLF)
        out.append(_row(("%-11s", "%-14s", "%-22s", "%-20s", "%-20s"),
                        ("Fiber No.", "Fiber Type", "Average PO2 (mmHg)", "StDev PO2 (mmHg)", "Hypoxic Fraction"))
                   + CRLF + CRLF)
        fm = ("%-11.0f", "%-14.0f", "%-22.2f", "%-20.3f", "%-20.4f")
        out.extend(_row(fm, r) + CRLF for r in _rows([fibers[c] for c in FIBER_PO2_COLUMNS]))
    return "".join(out)


def fiber_po2_table(solution: Any, fiber_index: Optional[Sequence[int]] = None,
                    fiber_types: Optional[Sequence[int]] = None) -> Dict[str, np.ndarray]:
    """``SingleFiberPO2Statistics``: one row per *geometry* fibre (1-based number,
    raw type code). ``fiber_index[j]`` is the geometry fibre of model fibre ``j``
    (``ModelGeometry.fiber_index``); fibres that are not in the model get NaN."""
    from .fem import fiber_po2_statistics

    st = fiber_po2_statistics(solution)
    nm = len(st["fiber"])
    idx = np.arange(nm) if fiber_index is None else np.asarray(fiber_index, int)
    n = len(fiber_types) if fiber_types is not None and len(fiber_types) else (int(idx.max()) + 1 if nm else 0)
    types = np.asarray(fiber_types, int) if fiber_types is not None and len(fiber_types) else np.zeros(n, int)
    if fiber_types is None or not len(fiber_types):
        types[idx] = st["fiber_type"]
    out = {"Fiber_No": np.arange(1, n + 1), "Fiber_Type": types}
    for src, dst in (("mean_po2_mmHg", "Average_PO2_mmHg"), ("std_po2_mmHg", "StDev_PO2_mmHg"),
                     ("hypoxic_fraction", "Hypoxic_Fraction")):
        col = np.full(n, np.nan)
        col[idx] = st[src]
        out[dst] = col
    return out


# ==========================================================================
# files
# ==========================================================================


def safe_base_name(name: str) -> str:
    """The user's file name without folder parts or characters Windows refuses."""
    name = os.path.basename(str(name).strip())
    for ch in '<>:"/\\|?*':
        name = name.replace(ch, "_")
    name = name.rstrip(" .")
    if not name:
        raise ValueError("empty file name")
    return name


def write_text(path: str, text: str) -> str:
    """Write ``text`` as-is (it already has CRLF line ends)."""
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(text)
    return path


def write_table_xlsx(path: str, columns: Mapping[str, Sequence[Any]], sheet: str = "Sheet1") -> str:
    """``writetable``: header row of column names, one row per entry; NaN -> empty."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    names = list(columns)
    ws.append(names)
    for c in ws[1]:
        c.font = Font(bold=True)
    cols = [np.asarray(columns[k]) for k in names]
    n = {len(c) for c in cols}
    if len(n) > 1:
        raise ValueError("columns of different length")
    for r in range(n.pop() if n else 0):
        row = []
        for c in cols:
            v = c[r]
            if isinstance(v, (np.integer, int)):
                row.append(int(v))
            else:
                v = float(v)
                row.append(v if math.isfinite(v) else None)
        ws.append(row)
    for k, name in enumerate(names, start=1):
        ws.column_dimensions[ws.cell(1, k).column_letter].width = max(10, len(name) + 2)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def read_table_xlsx(path: str) -> Dict[str, np.ndarray]:
    """Inverse of :func:`write_table_xlsx` (empty -> NaN)."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    names = list(rows[0])
    return {n: np.array([np.nan if r[k] is None else r[k] for r in rows[1:]], dtype=float)
            for k, n in enumerate(names)}


# ==========================================================================
# figures
# ==========================================================================


def _figure(size_in: Tuple[float, float]):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    fig = Figure(figsize=size_in, facecolor="white")
    FigureCanvasAgg(fig)
    return fig


def _save(fig: Any, path: str, dpi: int) -> str:
    fig.savefig(path, dpi=dpi, facecolor="white")
    return path


def histc(values: np.ndarray, edges: np.ndarray) -> np.ndarray:
    """MATLAB ``histc``: ``edges[k] <= v < edges[k+1]``; the last bin counts
    ``v == edges[-1]``; NaN and values outside are ignored."""
    v = np.asarray(values, float)
    v = v[np.isfinite(v)]
    e = np.asarray(edges, float)
    n = np.zeros(e.size)
    if e.size == 0 or v.size == 0:
        return n
    k = np.searchsorted(e, v, side="right") - 1
    ok = (k >= 0) & (k < e.size - 1)
    np.add.at(n, k[ok], 1)
    n[-1] += np.count_nonzero(v == e[-1])
    return n


def _index_bins(values: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    v = np.asarray(values, float)
    fin = v[np.isfinite(v)]
    nbars = max(1, int(round(math.sqrt(v.size) * 4)))
    vmax = float(fin.max()) if fin.size else 1.0
    if nbars == 1:
        x = np.array([vmax])
    else:
        x = np.linspace(0.0, vmax, nbars)
    return x, histc(v, x)


_TYPE_STYLE = ((22, "IIb", (1.0, 1.0, 1.0)), (21, "IIa", (0.75, 0.75, 0.75)), (1, "I", (0.5, 0.5, 0.5)))


def index_distribution_figure(values: np.ndarray, name: str, units: str = "",
                              fiber_types: Optional[np.ndarray] = None):
    """``saveIndexFigure`` panel: proportion and cumulative proportion histograms
    (``histc`` on ``round(4 sqrt(n))`` points from 0 to the maximum).

    For Fiber area, LCFR and LCD with at least two fibre types, the bars are split
    by type in 3-D (``bar3 ... 'detached'``) as in the legacy figure."""
    v = np.asarray(values, float)
    x, n = _index_bins(v)
    total = n.sum() or 1.0
    by_type = None
    if fiber_types is not None and name in ("LCD", "LCFR", "Fiber area"):
        t = np.asarray(fiber_types, int)
        present = [(code, lab, col) for code, lab, col in _TYPE_STYLE if np.any(t == code)]
        if len(present) >= 2 and t.size == v.size:
            by_type = [(lab, col, histc(v[t == code], x)) for code, lab, col in present]
            by_type.append(("All", (0.25, 0.25, 0.25), n))
    fig = _figure((11.0, 4.6))
    label = f"{name}{units}"
    font = {"family": "serif", "size": 10}
    if by_type is None:
        width = (x[1] - x[0]) if x.size > 1 else 1.0
        for k, (yvals, ylab) in enumerate(((n / total, "Proportion"), (np.cumsum(n) / total, "Cumulative proportion"))):
            ax = fig.add_subplot(1, 2, k + 1)
            ax.bar(x, yvals, width=width, color="black", edgecolor="white", linewidth=0.5, align="center")
            ax.set_xlabel(label, fontdict=font)
            ax.set_ylabel(ylab, fontdict=font)
            ax.set_xlim(0, max(x[-1] + width / 2, 1e-12))
            if k == 1:
                ax.set_ylim(0, 1.0)
            ax.set_box_aspect(1)
            ax.tick_params(labelsize=9)
    else:
        dy = (x[1] - x[0]) * 0.8 if x.size > 1 else 0.8
        for k, (cum, zlab) in enumerate(((False, "Proportion"), (True, "Cumulative proportion"))):
            ax = fig.add_subplot(1, 2, k + 1, projection="3d")
            for j, (lab, col, cnt) in enumerate(by_type):
                h = (np.cumsum(cnt) if cum else cnt) / total
                keep = h > 0
                ax.bar3d(np.full(keep.sum(), j + 1 - 0.35), x[keep] - dy / 2, np.zeros(keep.sum()),
                         0.7, dy, h[keep], color=col, edgecolor="k", linewidth=0.3, shade=True)
            ax.set_xticks(range(1, len(by_type) + 1))
            ax.set_xticklabels([b[0] for b in by_type])
            ax.set_xlabel("Fiber type", fontdict=font)
            ax.set_ylabel(label, fontdict=font)
            ax.set_zlabel(zlab, fontdict=font)
            ax.set_xlim(0.5, len(by_type) + 0.5)
            ax.set_ylim(max(x[-1], 1e-12), 0)                 # bar3 sets YDir 'reverse'
            ax.set_zlim(0, 1.05 if cum else float((n / total).max()) + 0.05)
            ax.view_init(elev=35, azim=-60 - 90)
            ax.tick_params(labelsize=8)
        fig.subplots_adjust(left=0.02, right=0.95, bottom=0.05, top=0.97, wspace=0.15)
        return fig
    fig.tight_layout()
    return fig


# ---- probability densities (PlotOxygenDistribution / PlotOxygenUptakeDistribution) -------------


def matlab_smooth(y: np.ndarray, span: float = 5) -> np.ndarray:
    """``smooth(y, span, 'moving')``: centred moving average whose window shrinks
    symmetrically at the ends; a fractional span is ``ceil(span * n)``; an even
    span is reduced by one."""
    y = np.asarray(y, float)
    n = y.size
    if n == 0:
        return y.copy()
    if span < 1:
        span = math.ceil(span * n)
    span = int(span)
    width = max(1, span - 1 + span % 2)
    h = width // 2
    out = np.empty(n)
    c = np.concatenate(([0.0], np.cumsum(y)))
    for i in range(n):
        r = min(h, i, n - 1 - i)
        out[i] = (c[i + r + 1] - c[i - r]) / (2 * r + 1)
    return out


def _binned_area(values: np.ndarray, area: np.ndarray, centres: np.ndarray, binsize: float) -> np.ndarray:
    """``sum(ar(u > x - b/2 & u <= x + b/2))`` for every centre ``x``."""
    edges = np.append(centres - binsize / 2, centres[-1] + binsize / 2)
    k = np.searchsorted(edges, values, side="left") - 1
    ok = (k >= 0) & (k < centres.size)
    return np.bincount(k[ok], weights=area[ok], minlength=centres.size)


@dataclass
class DensityCurves:
    x: np.ndarray
    pdf: Dict[str, np.ndarray]          # label -> probability density (normalised by the tissue area under)
    cdf: Dict[str, np.ndarray]          # label -> cumulative share of the total area
    xmin: float
    xlabel: str


def density_curves(values: np.ndarray, area: np.ndarray, compartment: Optional[np.ndarray], centres: np.ndarray,
                   binsize: float, xlabel: str, compartment_values: Optional[np.ndarray] = None) -> DensityCurves:
    """Legacy density/cumulative curves of a per-triangle field.

    ``compartment``: per-triangle code (0 IS, 1, 21, 22) or None (cardiac: tissue
    only). ``compartment_values`` replaces ``values`` for the compartment curves
    (the VO2 plot bins the interstitium at 0)."""
    from .fem.solve import COMPARTMENT_IS, FIBER_TYPE_I, FIBER_TYPE_IIA, FIBER_TYPE_IIB

    total = area.sum()
    a_all = _binned_area(values, area, centres, binsize)
    y_all = np.abs(matlab_smooth(a_all / total, 5))
    under = float(np.trapezoid(y_all, centres)) if hasattr(np, "trapezoid") else float(np.trapz(y_all, centres))
    under = under or 1.0
    pdf = {"Tissue": y_all / under}
    cdf = {"Tissue": np.cumsum(a_all / total)}
    if compartment is not None:
        cv = values if compartment_values is None else compartment_values
        for code, lab, span in ((COMPARTMENT_IS, "Interstitia", 5), (FIBER_TYPE_I, "Type I", 0.05),
                                (FIBER_TYPE_IIA, "Type IIa", 0.05), (FIBER_TYPE_IIB, "Type IIb", 0.05)):
            m = compartment == code
            if not m.any():
                continue
            a = _binned_area(cv[m], area[m], centres, binsize)
            pdf[lab] = np.abs(matlab_smooth(a / total, span)) / under
            cdf[lab] = np.cumsum(a / total)
    pos = np.zeros(centres.size, bool)
    for y in pdf.values():
        pos |= y > 0
    idx = int(np.argmax(pos)) + 1 if pos.any() else 1              # MATLAB 1-based IDmin
    if idx > 5:
        idx -= 4
    xmin = 0.0 if centres[idx - 1] <= 5 else float(math.floor(centres[idx - 1]))
    return DensityCurves(centres, pdf, cdf, xmin, xlabel)


_CURVE_COLOURS = {"Tissue": "k", "Interstitia": "c", "Type I": "r", "Type IIa": "g", "Type IIb": "b"}


def density_figure(curves: DensityCurves):
    """Two panels: probability density and cumulative probability."""
    fig = _figure((11.0, 4.4))
    for k, (data, ylab) in enumerate(((curves.pdf, "Probability density"),
                                      (curves.cdf, "Cumulative probability density"))):
        ax = fig.add_subplot(1, 2, k + 1)
        for lab, y in data.items():
            ax.plot(curves.x, y, color=_CURVE_COLOURS.get(lab, "k"), linewidth=1, label=lab)
        hi = curves.x[-1] if curves.x[-1] > curves.xmin else curves.xmin + 1
        pad = 0.005 * (hi - curves.xmin)               # keep a spike at the left edge visible
        ax.set_xlim(curves.xmin - pad, hi)
        if k == 1:
            ax.set_ylim(0, 1)
        else:
            ax.legend(loc="upper left", fontsize=8, frameon=True)
        ax.set_xlabel(curves.xlabel)
        ax.set_ylabel(ylab)
    fig.tight_layout()
    return fig


def _triangle_values(solution: Any) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """(u_nondim at centroids after the legacy clamp, areas ndim^2, compartment)."""
    from .fem.solve import _clamped

    mesh = solution.mesh
    u = _clamped(solution.u)[mesh.triangles].mean(axis=1)
    return u, mesh.cell_areas(), mesh.cell_compartment()


def po2_density(solution: Any, tissue: str = "skeletal") -> DensityCurves:
    """``PlotOxygenDistribution``: PO2 (mmHg), bins of P_c/2 from 0 to Pcap."""
    raw = solution.params.raw
    u, area, comp = _triangle_values(solution)
    U = raw.Pcap * u
    b = raw.P_c / 2
    centres = np.arange(0.0, raw.Pcap + b * 1e-9, b)
    return density_curves(U, area, None if tissue == "cardiac" else comp, centres, b, "Oxygen tension (mmHg)")


def _uptake_field(solution: Any, tissue: str) -> Tuple[np.ndarray, np.ndarray, np.ndarray, float]:
    """Per-triangle O2 consumption (ml/100 ml/min): ``6000 * VO2max_X * u/(u + pc)``,
    plus the saturation itself and ``M1`` (the scale of the legacy bins)."""
    u, area, comp = _triangle_values(solution)
    S = u / (u + solution.params.p_c)
    cf = solution.coefficients
    M0 = 6000.0 * solution.switches.exercise_level * solution.params.raw.M0
    if tissue == "cardiac":
        return M0 * S, S, area, M0
    M1 = cf.vol_avg_uptake_conversion * M0
    rate = np.zeros_like(S)
    for code, ratio in cf.uptake_ratio.items():
        rate[comp == code] = ratio * M1
    return rate * S, S, area, M1


def vo2_density(solution: Any, tissue: str = "skeletal") -> DensityCurves:
    """``PlotOxygenUptakeDistribution``: bins of M1 (max S - min S)/40 from 0 to M1."""
    V, S, area, M1 = _uptake_field(solution, tissue)
    rng = float(S.max() - S.min())
    b = M1 * rng / 40 if rng > 0 else max(M1, 1e-12) / 40
    nb = int(math.floor(M1 / b + 1e-9)) + 1
    if nb > 200000:                                  # a nearly uniform field: keep the curve drawable
        b = M1 / 200000
        nb = 200001
    centres = b * np.arange(nb)
    comp = None if tissue == "cardiac" else solution.mesh.cell_compartment()
    return density_curves(V, area, comp, centres, b, "Oxygen consumption (ml/100 ml/min)")


def _rings(f: Any) -> List[np.ndarray]:
    if hasattr(f, "exterior"):
        return [np.asarray(f.exterior.coords)]
    if hasattr(f, "geoms"):
        return [r for g in f.geoms for r in _rings(g)]
    if isinstance(f, (list, tuple)):
        return [np.asarray(r, float) for r in f]
    return [np.asarray(f, float)]


_OUTLINE = {1: (0, 0, 0), 21: (0.5, 0.5, 0.5), 22: (1, 1, 1)}


def _um(xy: np.ndarray, geometry: Any) -> np.ndarray:
    ls, frame = getattr(geometry, "length_scale", None), getattr(geometry, "frame", None)
    if ls is None or frame is None:
        return np.asarray(xy, float)
    return ls.ld * (np.asarray(xy, float) + np.asarray(frame.ss))


def profile_figure(solution: Any, geometry: Any, field: str = "po2", cmap: str = "jet",
                   clim: Optional[Tuple[float, float]] = None, tissue: str = "skeletal"):
    """``PO2_Profile`` / ``VO2_Profile``: the field over the tissue in um, fibre
    outlines by type (I black, IIa grey, IIb white, other green), colour bar below."""
    from matplotlib.tri import Triangulation

    mesh = solution.mesh
    pts = _um(mesh.points, geometry)
    tri = Triangulation(pts[:, 0], pts[:, 1], mesh.triangles)
    w = float(np.ptp(pts[:, 0])) or 1.0
    h = float(np.ptp(pts[:, 1])) or 1.0
    fig = _figure((7.0, 7.0 * h / w + 1.3))
    ax = fig.add_axes((0.03, 0.16, 0.94, 0.81))
    if field == "po2":
        vals = solution.params.raw.Pcap * np.clip(solution.u, 0, None)
        art = ax.tripcolor(tri, vals, shading="gouraud", cmap=cmap)
        label = "PO2 spatial profile (mmHg)"
    else:
        V, _, _, _ = _uptake_field(solution, tissue)
        art = ax.tripcolor(tri, facecolors=V, cmap=cmap)
        label = "VO2 spatial profile (ml/100 ml/min)"
    extend = "neither"
    if clim is not None:
        art.set_clim(*clim)
        if field == "po2":                                # fixed scale: arrows show values beyond it
            v = np.asarray(vals, float)
            lo, hi = v.min() < clim[0], v.max() > clim[1]
            extend = "both" if lo and hi else "min" if lo else "max" if hi else "neither"
            label = f"PO2 spatial profile (mmHg), fixed scale {clim[0]:g}-{clim[1]:g}"
    if tissue != "cardiac":
        for f, t in zip(getattr(geometry, "fibers", []), getattr(geometry, "fiber_types", [])):
            for r in _rings(f):
                p = _um(r, geometry)
                ax.plot(np.append(p[:, 0], p[0, 0]), np.append(p[:, 1], p[0, 1]),
                        color=_OUTLINE.get(int(t), (0, 0.8, 0)), linewidth=0.8)
    ax.set_xlim(pts[:, 0].min(), pts[:, 0].max())
    ax.set_ylim(pts[:, 1].min(), pts[:, 1].max())
    ax.set_aspect("equal")
    ax.axis("off")
    cax = fig.add_axes((0.15, 0.08, 0.7, 0.03))
    cb = fig.colorbar(art, cax=cax, orientation="horizontal", extend=extend)
    cb.set_label(label, fontsize=12, family="serif")
    return fig


_FILL = {1: (0.35, 0.35, 0.35), 21: (0.0, 0.9, 0.9), 22: (0.9, 0.0, 0.9)}


def flux_figure(flux: Any, geometry: Any, mesh: Any, voronoi: Any = None, roi: Any = None):
    """``OxygenFluxGUI``: fibres shaded by type (I dark grey, IIa cyan, IIb magenta,
    unknown green), capillary domains in red, capillaries as red disks and the
    flux lines in black; dashed ROI rectangle."""
    from matplotlib.collections import LineCollection, PatchCollection
    from matplotlib.patches import Circle, Polygon, Rectangle

    pts = _um(mesh.points, geometry)
    x0, x1 = pts[:, 0].min(), pts[:, 0].max()
    y0, y1 = pts[:, 1].min(), pts[:, 1].max()
    fig = _figure((8.0, 8.0 * (y1 - y0) / max(x1 - x0, 1e-12) + 0.2))
    ax = fig.add_axes((0.01, 0.01, 0.98, 0.98))
    patches, colours = [], []
    for f, t in zip(getattr(geometry, "fibers", []), getattr(geometry, "fiber_types", [])):
        for r in _rings(f):
            patches.append(Polygon(_um(r, geometry), closed=True))
            colours.append(_FILL.get(int(t), (0.0, 0.8, 0.0)))
    if patches:
        ax.add_collection(PatchCollection(patches, facecolors=colours, edgecolors="k", linewidths=0.4, alpha=0.6))
    if voronoi is not None:
        segs = []
        for c in voronoi.cells:
            if c is None or getattr(c, "is_empty", True):
                continue
            for r in _rings(c):
                segs.append(_um(r, geometry))
        ax.add_collection(LineCollection(segs, colors="r", linewidths=0.8))
    ld = geometry.length_scale.ld if getattr(geometry, "length_scale", None) is not None else 1.0
    caps = _um(mesh.capillaries, geometry)
    ax.add_collection(PatchCollection([Circle(c, r * ld) for c, r in zip(caps, mesh.rcap)], facecolors="r",
                                      edgecolors="r"))
    lines = []
    for i, group in enumerate(_flux_paths(flux)):
        for p in group:
            if len(p) > 1:
                lines.append(_um(p, geometry))
    ax.add_collection(LineCollection(lines, colors="k", linewidths=0.8))
    if roi is not None:
        ax.add_patch(Rectangle((roi.x_min_um, roi.y_min_um), roi.x_max_um - roi.x_min_um,
                               roi.y_max_um - roi.y_min_um, fill=False, edgecolor="b", linestyle="--", linewidth=1))
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect("equal")
    ax.axis("off")
    return fig


def _dedupe(p: np.ndarray) -> np.ndarray:
    p = np.asarray(p, float)
    if len(p) < 2:
        return p
    keep = np.r_[True, np.any(np.diff(p, axis=0) != 0, axis=1)]
    return p[keep]


def _flux_paths(flux: Any) -> List[List[np.ndarray]]:
    paths = getattr(flux, "paths", None)
    st = np.asarray(flux.streams)
    return [[_dedupe(paths[i][j] if paths else st[i, j]) for j in range(st.shape[1])] for i in range(st.shape[0])]


def write_flux_csv(path: str, flux: Any, geometry: Any, mesh: Any,
                   capillary_numbers: Optional[Sequence[int]] = None) -> str:
    """One row per point: line, capillary, seed, point, x_um, y_um, s_um (arc length
    from the seed), stop_reason. ``capillary`` is the 1-based number of the
    capillary in the data file when ``capillary_numbers`` (model -> data index,
    0-based) is given, otherwise the 1-based model capillary."""
    ld = geometry.length_scale.ld if getattr(geometry, "length_scale", None) is not None else 1.0
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["line", "capillary", "capillary_x_um", "capillary_y_um", "seed", "point", "x_um", "y_um", "s_um",
                    "stop_reason"])
        line = 0
        caps = _um(mesh.capillaries, geometry)
        for i, group in enumerate(_flux_paths(flux)):
            mc = int(flux.capillary_index[i])
            num = (int(capillary_numbers[mc]) + 1) if capillary_numbers is not None else mc + 1
            for j, p in enumerate(group):
                if len(p) < 2:
                    continue
                line += 1
                q = _um(p, geometry)
                s = np.r_[0.0, np.cumsum(np.hypot(*np.diff(p, axis=0).T))] * ld
                why = str(flux.stop_reason[i, j])
                for k in range(len(q)):
                    w.writerow([line, num, f"{caps[mc, 0]:.4f}", f"{caps[mc, 1]:.4f}", j + 1, k + 1,
                                f"{q[k, 0]:.6f}", f"{q[k, 1]:.6f}", f"{s[k]:.6f}", why])
    return path


# ==========================================================================
# high level
# ==========================================================================


def _prep(folder: str, base: str) -> str:
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, safe_base_name(base))


def _style_and_dpi(style: Any, dpi: Optional[int]):
    """``style`` None: the legacy MATLAB figures as PNG (300 dpi unless ``dpi``)."""
    from .figures import FigureStyle

    if style is None:
        style = FigureStyle.legacy(dpi=int(dpi or 300))
    return style, int(dpi or style.dpi)


def _write(fig: Any, stem: str, style: Any, dpi: int) -> List[str]:
    from .figures import save_figure

    return save_figure(fig, stem, style, dpi)


def export_indices(morphometrics: Any, geometry: Any, folder: str, base: str, figures: bool = True,
                   dpi: Optional[int] = None, progress: Optional[ProgressCallback] = None,
                   style: Any = None) -> List[str]:
    """``IndicesMenu`` Save: text summary, per-fibre spreadsheet and distribution
    figures in ``folder`` (normally ``<data dir>/INDICES``). ``style``
    (:class:`otm_core.figures.FigureStyle`): figure layout and file formats;
    None = the legacy MATLAB figures as PNG."""
    style, dpi = _style_and_dpi(style, dpi)
    stem = _prep(folder, base)
    rep = IndicesReport.from_morphometrics(morphometrics, getattr(geometry, "fiber_types", None))
    out = [write_text(stem + "_Output_Text_File.txt", format_indices_text(rep))]
    report(progress, 0.05, "text file")
    if rep.has_fibres:
        out.append(write_table_xlsx(stem + "_FiberIndicesOutput.xlsx", rep.fiber_table(), "FiberIndices"))
    if figures:
        panels = rep.distributions()
        for k, (name, units, values) in enumerate(panels):
            report(progress, 0.1 + 0.9 * k / len(panels), f"figure: {name}")
            types = rep.fiber_type if rep.has_fibres and values.size == rep.fiber_type.size else None
            if style.publication:
                from .figures import index_figure

                fig = index_figure(values, name, units, types, style)
            else:
                fig = index_distribution_figure(values, name, units, types)
            out += _write(fig, f"{stem}_{name}_Distribution", style, dpi)
    report(progress, 1.0, "done")
    return out


def export_po2(solution: Any, geometry: Any, folder: str, base: str, tissue: str = "skeletal",
               fiber_index: Optional[Sequence[int]] = None, figures: bool = True, cmap: str = "jet",
               clim: Optional[Tuple[float, float]] = None, dpi: Optional[int] = None,
               progress: Optional[ProgressCallback] = None, style: Any = None) -> List[str]:
    """``FEMSimulation`` Save: PO2/MO2 statistics (text), per-fibre PO2
    (spreadsheet, skeletal) and the PO2/VO2 profile and density figures in
    ``folder`` (normally ``<data dir>/PO2``)."""
    from .fem import mo2_statistics, po2_statistics

    style, dpi = _style_and_dpi(style, dpi)
    stem = _prep(folder, base)
    po2, mo2 = po2_statistics(solution), mo2_statistics(solution)
    out: List[str] = []
    if tissue == "cardiac":
        out.append(write_text(stem + "_PO2_Stats_Output_Text_File.txt", format_po2_text(po2, mo2, tissue="cardiac")))
    else:
        fib = fiber_po2_table(solution, fiber_index, getattr(geometry, "fiber_types", None))
        out.append(write_text(stem + "_PO2_Stats.txt", format_po2_text(po2, mo2, fib)))
        out.append(write_table_xlsx(stem + "_FiberPO2Stats.xlsx", fib, "FiberPO2"))
    report(progress, 0.1, "statistics")
    if figures:
        if style.publication:
            from . import figures as pf

            steps = (("PO2_Profile", lambda: pf.map_figure(solution, geometry, "po2", style, cmap, clim, tissue)),
                     ("PO2_ProbabilityDensity",
                      lambda: pf.po2_distribution_figure(po2_density(solution, tissue), style)),
                     ("VO2_ProbabilityDensity", lambda: pf.vo2_figure(solution, tissue, style)),
                     ("VO2_Profile", lambda: pf.map_figure(solution, geometry, "vo2", style, cmap, None, tissue)))
        else:
            steps = (("PO2_Profile", lambda: profile_figure(solution, geometry, "po2", cmap, clim, tissue)),
                     ("PO2_ProbabilityDensity", lambda: density_figure(po2_density(solution, tissue))),
                     ("VO2_ProbabilityDensity", lambda: density_figure(vo2_density(solution, tissue))),
                     ("VO2_Profile", lambda: profile_figure(solution, geometry, "vo2", cmap, None, tissue)))
        for k, (name, make) in enumerate(steps):
            report(progress, 0.1 + 0.9 * k / len(steps), f"figure: {name}")
            out += _write(make(), f"{stem}_{name}", style, dpi)
    report(progress, 1.0, "done")
    return out


def export_flux(flux: Any, geometry: Any, mesh: Any, folder: str, base: str, voronoi: Any = None,
                capillary_numbers: Optional[Sequence[int]] = None, figure: bool = True,
                dpi: Optional[int] = None, progress: Optional[ProgressCallback] = None,
                style: Any = None) -> List[str]:
    """``OxygenFluxGUI`` Save (figure) plus the line coordinates as CSV."""
    style, dpi = _style_and_dpi(style, dpi)
    stem = _prep(folder, base)
    out = [write_flux_csv(stem + "_PO2_Flux_Lines.csv", flux, geometry, mesh, capillary_numbers)]
    report(progress, 0.3, "coordinates")
    if figure:
        if style.publication:
            from .figures import flux_map_figure

            fig = flux_map_figure(flux, geometry, mesh, getattr(geometry, "roi", None), style)
        else:
            fig = flux_figure(flux, geometry, mesh, voronoi, getattr(geometry, "roi", None))
        out += _write(fig, stem + "_PO2_Flux_Lines", style, dpi)
    report(progress, 1.0, "done")
    return out
