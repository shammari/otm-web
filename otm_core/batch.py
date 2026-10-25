"""A whole experiment at once: a folder tree of Dtect exports with grouped summaries.

Port of the OTM 2026 (MATLAB) batch (``otm.batch.run``). The folders between the
experiment's top folder and each file are its groups, e.g.::

    TA study/Animal 3/Region B/s01.mat   ->  Animal = "Animal 3", Region = "Region B"

with ``level_names=("Animal", "Region")``. :func:`run_experiment` runs every file
with one :class:`~otm_core.pipeline.RunSettings` (failed files are recorded and
the batch goes on) and writes to the output folder:

* ``OTM_batch_summary.xlsx`` with the sheets
  - *Summary*: one row per file, with a column per folder level;
  - *By <all levels>*: mean and SD of every column per group (samples);
  - *By <levels 2..> (across <level 1>)*: the samples of each unit (e.g. each
    animal) averaged first, then mean, SD and SEM over the units;
  - *By <level 1>*: e.g. each animal, regions pooled;
  - *Fibres*: one row per fibre of every sample (size, indices, PO2, VO2);
  - *Capillary domains*: one row per capillary domain (area, PO2, VO2, % hypoxia);
  - *Data checks*: every problem the data check found and what was done;
* the same tables as CSV files (``OTM_batch_*.csv``);
* ``OTM_batch_log.txt`` (outcome of every file) and ``OTM_batch_settings.json``.
"""

from __future__ import annotations

import csv
import math
import os
import re
import time
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Mapping, Optional, Sequence, Tuple

import numpy as np

from .progress import ProgressCallback, report

__all__ = ["read_size_table", "GroupInfo", "group_info", "group_stats", "natural_key", "fibre_rows", "domain_rows",
           "data_check_rows", "find_experiment_files", "run_experiment", "ExperimentResult", "write_workbook",
           "LOG_FILE"]

LOG_FILE = "OTM_batch_log.txt"
_TYPE_NAMES = {1: "I", 21: "IIa", 22: "IIb", 0: "unknown"}
_RESERVED = {"file", "sample", "folder", "status", "errors", "tissue", "exercise"}


# ==========================================================================
# groups from the folder tree
# ==========================================================================


@dataclass
class GroupInfo:
    root: str
    names: List[str]                       # level names, e.g. ["Animal", "Region"]
    columns: List[str]                     # column names for the tables (valid, unique)
    levels: List[List[str]]                # per file, one value per level ("" where a file sits higher)
    samples: List[str]                     # per file, "Animal 3/Region B/s01"
    folders: List[List[str]]               # per file, the folders between root and the file


def _split(p: str) -> List[str]:
    p = os.path.normpath(os.path.abspath(p))
    drive, rest = os.path.splitdrive(p)
    parts = [x for x in re.split(r"[\\/]+", rest) if x]
    return ([drive] if drive else []) + parts


def _same(a: str, b: str) -> bool:
    return a.lower() == b.lower() if os.name == "nt" else a == b


def group_info(files: Sequence[str], root: Optional[str] = None,
               level_names: Sequence[str] | str = ()) -> GroupInfo:
    """Groups of ``files`` from their folders below ``root`` (default: the deepest
    folder holding every file). Missing level names become "Level 1", "Level 2" ..."""
    if isinstance(level_names, str):
        level_names = [s.strip() for s in level_names.split(",")]
    level_names = [s.strip() for s in level_names if s and s.strip()]
    parts = [_split(os.path.dirname(f)) for f in files]
    stems = [os.path.splitext(os.path.basename(f))[0] for f in files]

    def common(ps):
        if not ps:
            return []
        out = []
        for col in zip(*ps):
            if all(_same(col[0], c) for c in col):
                out.append(col[0])
            else:
                break
        return out

    rp = common(parts)
    if root:
        r = _split(root)
        if all(len(p) >= len(r) and all(_same(a, b) for a, b in zip(p, r)) for p in parts):
            rp = r
    root_s = os.path.join(*rp) if rp else ""
    if rp and rp[0].endswith(":"):
        root_s = rp[0] + os.sep + os.path.join(*rp[1:]) if len(rp) > 1 else rp[0] + os.sep
    elif rp and not os.path.splitdrive(rp[0])[0]:
        root_s = os.sep + os.path.join(*rp)
    folders = [p[len(rp):] for p in parts]
    K = max((len(f) for f in folders), default=0)
    levels = [f + [""] * (K - len(f)) for f in folders]
    names = [level_names[k] if k < len(level_names) else f"Level {k + 1}" for k in range(K)]
    cols: List[str] = []
    for nm in names:
        c = re.sub(r"[^0-9A-Za-z_]+", "_", nm).strip("_") or "level"
        if c[0].isdigit():
            c = "L" + c
        if c.lower() in _RESERVED:
            c = c + "_group"
        base, j = c, 2
        while c in cols:
            c, j = f"{base}_{j}", j + 1
        cols.append(c)
    samples = ["/".join(f + [s]) for f, s in zip(folders, stems)]
    return GroupInfo(root_s, names, cols, levels, samples, folders)


def find_experiment_files(folders: Sequence[str], pattern: str = "*.mat") -> List[str]:
    """Every ``*.mat`` below the folders (sub-folders searched), sorted naturally;
    files the batch itself writes (sessions, reports) are skipped."""
    import fnmatch

    out: List[str] = []
    for top in folders:
        if os.path.isfile(top):
            out.append(os.path.abspath(top))
            continue
        for dp, dn, fn in os.walk(top):
            dn[:] = [d for d in dn if d not in ("INDICES", "PO2") and not d.startswith(".")]
            for f in fn:
                if fnmatch.fnmatch(f.lower(), pattern.lower()) and not f.startswith("~$"):
                    out.append(os.path.abspath(os.path.join(dp, f)))
    seen, res = set(), []
    for f in sorted(out, key=lambda p: [natural_key(x) for x in _split(p)]):
        if f not in seen:
            seen.add(f)
            res.append(f)
    return res


def read_size_table(path: str) -> Dict[str, Dict[str, Any]]:
    """Per-file sizes from a CSV or Excel table with a ``file`` column and
    ``width_um`` or ``height_um`` (optionally ``roi_x_min_um`` .. ``roi_y_max_um``)."""
    from .pipeline import read_dimensions_table

    if not path.lower().endswith((".xlsx", ".xlsm")):
        return read_dimensions_table(path)
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return {}
    import io as _io

    buf = _io.StringIO()
    w = csv.writer(buf)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    tmp = path + ".csv.tmp"
    try:
        with open(tmp, "w", encoding="utf-8", newline="") as fh:
            fh.write(buf.getvalue())
        return read_dimensions_table(tmp)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


# ==========================================================================
# grouped statistics
# ==========================================================================


def natural_key(s: Any) -> Tuple:
    """Sort key with numbers compared as numbers ("Animal 2" before "Animal 10")."""
    out = []
    for tok in re.findall(r"\d+|\D+", str(s).lower()):
        out.append((0, int(tok), "") if tok.isdigit() else (1, 0, tok))
    return tuple(out)


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool)


def _msd(v: List[float]) -> Tuple[float, float]:
    a = np.asarray([x for x in v if x is not None and _is_num(x) and math.isfinite(x)], float)
    if a.size == 0:
        return math.nan, math.nan
    return float(a.mean()), float(a.std(ddof=1)) if a.size > 1 else math.nan


def group_stats(rows: Sequence[Mapping[str, Any]], keys: Sequence[str], unit_key: str = "",
                count_name: str = "n_units") -> List[Dict[str, Any]]:
    """Mean and SD of every numeric column per group of ``keys``.

    With ``unit_key`` (e.g. "Animal"): the samples are averaged within each unit
    first; the unit means are then summarised per group (mean, SD, SEM = SD /
    sqrt(units)). Groups come in natural order."""
    rows = list(rows)
    if not rows:
        return []
    keys = list(keys)
    skip = set(keys) | ({unit_key} if unit_key else set())
    cols: List[str] = []
    for r in rows:
        cols += [c for c in r if c not in cols and c not in skip]
    num = [c for c in cols if all(_is_num(r.get(c, math.nan)) for r in rows)]

    def groups(rs, ks):
        g: Dict[Tuple, List] = {}
        for r in rs:
            g.setdefault(tuple(str(r.get(k, "")) for k in ks), []).append(r)
        return sorted(g.items(), key=lambda kv: [natural_key(x) for x in kv[0]])

    out: List[Dict[str, Any]] = []
    if not unit_key:
        for kv, sel in groups(rows, keys):
            r = dict(zip(keys, kv))
            r["n_samples"] = len(sel)
            for c in num:
                m, s = _msd([x.get(c, math.nan) for x in sel])
                r[f"{c}_mean"], r[f"{c}_sd"] = m, s
            out.append(r)
        return out
    units = []
    for kv, sel in groups(rows, [unit_key] + keys):
        u = dict(zip([unit_key] + keys, kv))
        u["n_samples"] = len(sel)
        for c in num:
            u[c] = _msd([x.get(c, math.nan) for x in sel])[0]
        units.append(u)
    for kv, sel in groups(units, keys):
        r = dict(zip(keys, kv))
        r[count_name] = len(sel)
        r["n_samples"] = int(sum(u["n_samples"] for u in sel))
        for c in num:
            v = [u[c] for u in sel]
            m, s = _msd(v)
            n = sum(1 for x in v if math.isfinite(x))
            r[f"{c}_mean"], r[f"{c}_sd"], r[f"{c}_sem"] = m, s, (s / math.sqrt(n) if n else math.nan)
        out.append(r)
    return out


# ==========================================================================
# pooled tables (computed from the full result, before its arrays are dropped)
# ==========================================================================


def _ring_area_perimeter(f: Any) -> Tuple[float, float, np.ndarray]:
    from .export import _rings

    rings = _rings(f)
    if not rings:
        return math.nan, math.nan, np.full(2, np.nan)
    r = np.asarray(rings[0], float)
    x, y = r[:, 0], r[:, 1]
    a = 0.5 * abs(np.sum(x * np.roll(y, -1) - np.roll(x, -1) * y))
    p = float(np.sum(np.hypot(np.diff(np.r_[x, x[:1]]), np.diff(np.r_[y, y[:1]]))))
    for h in rings[1:]:
        h = np.asarray(h, float)
        a -= 0.5 * abs(np.sum(h[:, 0] * np.roll(h[:, 1], -1) - np.roll(h[:, 0], -1) * h[:, 1]))
    return a, p, r.mean(axis=0)


def _triangle_fields(res: Any):
    """Per mesh triangle: PO2 (mmHg), VO2 (ml/100 ml/min), area (ndim^2), or None."""
    sol = res.solution
    if sol is None:
        return None
    from .export import _uptake_field
    from .fem.solve import _clamped

    mesh = sol.mesh
    po2 = sol.params.raw.Pcap * _clamped(sol.u)[mesh.triangles].mean(axis=1)
    vo2, _s, area, _m = _uptake_field(sol, res.settings.tissue)
    return po2, vo2, area, sol.params.raw.P_c


def _area_stats(vals: np.ndarray, area: np.ndarray, pc: float) -> Tuple[float, float, float, float]:
    if area.sum() <= 0:
        return (math.nan,) * 4
    m = float(np.sum(vals * area) / area.sum())
    sd = float(math.sqrt(np.sum((vals - m) ** 2 * area) / area.sum()))
    return m, sd, float(vals.min()), float(100.0 * area[vals <= pc].sum() / area.sum())


def fibre_rows(res: Any) -> List[Dict[str, Any]]:
    """One row per fibre of a run (``RunResult`` with its arrays): type, size, the
    supply indices (ROI fibres), PO2 and O2 consumption (modelled fibres)."""
    g, m = res.geometry, res.morphometrics
    if g is None or not len(g.fibers) or g.length_scale is None or res.settings.tissue == "cardiac":
        return []
    n = len(g.fibers)
    ls = g.length_scale
    h = min(ls.x_um, ls.y_um) / 2.0                                 # um per model unit (otm_core.models.ROI.to_ndim)
    nan = math.nan
    rows = []
    types = np.asarray(g.fiber_types, int) if len(g.fiber_types) == n else np.zeros(n, int)
    for k, f in enumerate(g.fibers):
        a, p, c = _ring_area_perimeter(f)
        rows.append({"fibre": k + 1, "type": int(types[k]), "type_name": _TYPE_NAMES.get(int(types[k]), str(types[k])),
                     "in_roi": 0, "centroid_x_um": float(c[0] * h + ls.x_um / 2) if np.isfinite(c[0]) else nan,
                     "centroid_y_um": float(c[1] * h + ls.y_um / 2) if np.isfinite(c[1]) else nan,
                     "area_um2": a * h * h, "perimeter_um": p * h})
    for r in rows:
        for col in ("LCFR", "LCD_per_mm2", "DFR", "CFi", "CFPE", "Dmax_um"):
            r[col] = nan
    s = m.supply if m is not None else None
    if s is not None:
        idx = np.asarray(s.roi_fiber_index, int)
        cols = (("LCFR", s.LCFR), ("LCD_per_mm2", s.LCD), ("DFR", s.DFR), ("CFi", s.CFi), ("CFPE", s.CFPE),
                ("Dmax_um", s.Dmax))
        for j, k in enumerate(idx):
            if 0 <= k < n:
                rows[k]["in_roi"] = 1
                for name, st in cols:
                    v = np.asarray(st.values, float)
                    rows[k][name] = float(v[j]) if j < v.size else nan
    tf = _triangle_fields(res)
    for r in rows:
        r.update({"po2_mean_mmHg": nan, "po2_sd_mmHg": nan, "po2_min_mmHg": nan, "po2_hypoxic_pct": nan,
                  "vo2_mean_ml_100ml_min": nan})
    if tf is not None:
        po2, vo2, area, pc = tf
        mesh = res.solution.mesh
        reg = np.asarray(mesh.cell_region)
        fidx = getattr(res.model, "fiber_index", None)
        fidx = np.arange(len(mesh.fiber_types)) if fidx is None else np.asarray(fidx, int)
        for j, k in enumerate(fidx):
            sel = reg == j
            if 0 <= k < n and sel.any():
                mm, sd, mn, hyp = _area_stats(po2[sel], area[sel], pc)
                rows[k].update({"po2_mean_mmHg": mm, "po2_sd_mmHg": sd, "po2_min_mmHg": mn, "po2_hypoxic_pct": hyp,
                                "vo2_mean_ml_100ml_min": float(np.sum(vo2[sel] * area[sel]) / area[sel].sum())})
    return rows


def domain_rows(res: Any) -> List[Dict[str, Any]]:
    """One row per capillary domain (Voronoi cell) of a run: size, the ROI indices,
    and the PO2 / O2 consumption of the modelled tissue closest to the capillary."""
    g, m = res.geometry, res.morphometrics
    if g is None or not len(g.capillaries) or g.length_scale is None:
        return []
    caps = np.asarray(g.capillaries, float)
    n = len(caps)
    ls = g.length_scale
    h = min(ls.x_um, ls.y_um) / 2.0
    nan = math.nan
    rows = [{"capillary": k + 1, "x_um": float(caps[k, 0] * h + ls.x_um / 2),
             "y_um": float(caps[k, 1] * h + ls.y_um / 2), "in_roi": 0, "domain_area_um2": nan, "eq_diameter_um": nan,
             "FDR": nan, "SF": nan, "nn_mean_um": nan, "nn_min_um": nan} for k in range(n)]
    if m is not None:
        d = m.domains
        areas = np.asarray(d.capillary_domain_areas_um2, float)
        for k in range(min(n, areas.size)):
            a = areas[k]
            rows[k]["domain_area_um2"] = float(a)
            rows[k]["eq_diameter_um"] = float(2 * math.sqrt(a / math.pi)) if np.isfinite(a) and a > 0 else nan
        roi = np.asarray(d.roi_index, int)
        s, nn = m.supply, m.nearest_neighbours
        fdr = np.asarray(s.FDR.values, float) if s is not None else np.zeros(0)
        sf = np.asarray(s.SF.values, float) if s is not None else np.zeros(0)
        means = np.asarray(nn.means, float) if nn is not None else np.zeros(0)
        mins = np.asarray(nn.mins, float) if nn is not None else np.zeros(0)
        for j, k in enumerate(roi):
            if 0 <= k < n:
                rows[k]["in_roi"] = 1
                rows[k]["FDR"] = float(fdr[j]) if j < fdr.size else nan
                rows[k]["SF"] = float(sf[j]) if j < sf.size else nan
                rows[k]["nn_mean_um"] = float(means[j]) if j < means.size else nan
                rows[k]["nn_min_um"] = float(mins[j]) if j < mins.size else nan
    for r in rows:
        r.update({"tissue_area_um2": nan, "po2_mean_mmHg": nan, "po2_sd_mmHg": nan, "po2_min_mmHg": nan,
                  "po2_hypoxic_pct": nan, "vo2_mean_ml_100ml_min": nan})
    tf = _triangle_fields(res)
    if tf is not None:
        from scipy.spatial import cKDTree

        po2, vo2, area, pc = tf
        mesh = res.solution.mesh
        cen = np.asarray(mesh.points, float)[mesh.triangles].mean(axis=1)
        owner = cKDTree(caps).query(cen)[1]                         # the Voronoi cell holding the centre
        for k in range(n):
            sel = owner == k
            if sel.any():
                mm, sd, mn, hyp = _area_stats(po2[sel], area[sel], pc)
                rows[k].update({"tissue_area_um2": float(area[sel].sum() * h * h), "po2_mean_mmHg": mm,
                                "po2_sd_mmHg": sd, "po2_min_mmHg": mn, "po2_hypoxic_pct": hyp,
                                "vo2_mean_ml_100ml_min": float(np.sum(vo2[sel] * area[sel]) / area[sel].sum())})
    return rows


def data_check_rows(check: Optional[Mapping[str, Any]], status: str = "") -> List[Dict[str, Any]]:
    """One row per problem the data check found in a file, with what was done."""
    if not check:
        return []
    acts = check.get("actions", [])
    removed = {c for a in acts if a.get("action") == "removed" for c in a.get("codes", [])}
    kept = {c for a in acts if a.get("action") == "kept" for c in a.get("codes", [])}
    out = []
    for it in list(check.get("fixed_issues", [])) + list(check.get("issues", [])):
        c = it["code"]
        if c in removed and it in check.get("fixed_issues", []):
            done = "removed by OTM"
        elif c in kept:
            done = "kept by the user"
        elif it["severity"] in ("error", "fatal") and status != "ok":
            done = "file not analysed"
        else:
            done = "warning only"
        out.append({"code": c, "severity": it["severity"], "problem": it["title"], "message": it["message"],
                    "items": "; ".join(it.get("items", [])[:20]) + (" ..." if len(it.get("items", [])) > 20 else ""),
                    "action": done})
    return out


# ==========================================================================
# writing
# ==========================================================================


def _cell(v: Any) -> Any:
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return float(v) if math.isfinite(v) else None
    return v


def _csv_text(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, (float, np.floating)):
        return "" if not math.isfinite(v) else repr(float(v))
    return v


def _columns(rows: Sequence[Mapping[str, Any]]) -> List[str]:
    cols: List[str] = []
    for r in rows:
        cols += [c for c in r if c not in cols]
    # drop columns that are empty in every row
    return [c for c in cols if any(_cell(r.get(c)) not in (None, "") for r in rows)] or cols


def _sheet_name(s: str, used: List[str]) -> str:
    s = re.sub(r"[:\\/?*\[\]]", "-", s)[:31]
    base, j = s, 2
    while s.lower() in (u.lower() for u in used):
        suffix = f" ({j})"
        s, j = base[:31 - len(suffix)] + suffix, j + 1
    used.append(s)
    return s


def write_workbook(sheets: Sequence[Tuple[str, str, Sequence[Mapping[str, Any]]]], folder: str,
                   workbook: str = "OTM_batch_summary.xlsx") -> List[str]:
    """``sheets`` = (sheet name, csv stem, rows). Writes one CSV per sheet and one
    Excel workbook with every sheet. Returns the files written."""
    os.makedirs(folder, exist_ok=True)
    files: List[str] = []
    for _, stem, rows in sheets:
        if not rows:
            continue
        cols = _columns(rows)
        p = os.path.join(folder, stem + ".csv")
        with open(p, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            for r in rows:
                w.writerow([_csv_text(r.get(c)) for c in cols])
        files.append(p)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:                                            # pragma: no cover
        return files
    wb = Workbook()
    wb.remove(wb.active)
    used: List[str] = []
    for name, _, rows in sheets:
        if not rows:
            continue
        cols = _columns(rows)
        ws = wb.create_sheet(_sheet_name(name, used))
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_cell(r.get(c)) for c in cols])
        ws.freeze_panes = "B2"
        for k, nm in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, k).column_letter].width = min(40, max(9, len(str(nm)) + 2))
    if used:
        p = os.path.join(folder, workbook)
        wb.save(p)
        files.append(p)
    return files


# ==========================================================================
# running an experiment
# ==========================================================================


@dataclass
class ExperimentResult:
    files: List[str]
    groups: GroupInfo
    results: List[Any]                     # RunResult per file (light: no large arrays)
    rows: List[Dict[str, Any]]             # summary rows (with the group columns)
    written: List[str] = field(default_factory=list)
    out_dir: str = ""
    seconds: float = 0.0
    cancelled: bool = False

    @property
    def n_ok(self) -> int:
        return sum(1 for r in self.rows if r.get("status") == "ok")

    @property
    def n_failed(self) -> int:
        return sum(1 for r in self.rows if r.get("status") not in ("ok", "not run"))


def run_experiment(files: Sequence[str], settings: Any = None, root: Optional[str] = None,
                   level_names: Sequence[str] | str = (), out_dir: Optional[str] = None,
                   dims: Optional[Mapping[str, Mapping[str, Any]]] = None, save_sessions: bool = False,
                   apply_saved_edits: bool = False, group_sheets: bool = True, pooled_tables: bool = True,
                   progress: Optional[ProgressCallback] = None,
                   on_result: Optional[Callable[[int, int, Any], None]] = None,
                   cancelled: Optional[Callable[[], bool]] = None) -> ExperimentResult:
    """Run every file and write the summary workbook (see the module docstring).

    Each file writes its usual outputs: ``INDICES``/``PO2`` next to the data file,
    or under ``out_dir/<the same sub-folders>/<file name>/``. ``cancelled()`` is
    asked between files; the summary is written for the files that ran."""
    from .pipeline import RunResult, RunSettings, SETTINGS_FILE, _settings_for, find_saved_edits, run_file, \
        summary_row

    s = settings or RunSettings()
    files = [os.path.abspath(f) for f in files]
    gi = group_info(files, root, level_names)
    out = out_dir or gi.root or (os.path.dirname(files[0]) if files else os.getcwd())
    os.makedirs(out, exist_ok=True)
    t0 = time.perf_counter()
    n = len(files)
    results: List[Any] = []
    rows: List[Dict[str, Any]] = []
    fib: List[Dict[str, Any]] = []
    dom: List[Dict[str, Any]] = []
    checks: List[Dict[str, Any]] = []
    log = [f"OTM batch, {time.strftime('%Y-%m-%d %H:%M:%S')}: {n} file(s) under {gi.root or out}"]
    stop = False
    for i, path in enumerate(files):
        prefix = {"file": os.path.basename(path)}
        prefix.update({c: v for c, v in zip(gi.columns, gi.levels[i])})
        prefix["sample"] = gi.samples[i]
        if stop or (cancelled is not None and cancelled()):
            stop = True
            rows.append({**prefix, "status": "not run", "errors": "batch cancelled"})
            log.append(f"[{i + 1}/{n}] {gi.samples[i]}: not run (cancelled)")
            continue
        report(progress, i / max(n, 1), f"File {i + 1} of {n}: {os.path.basename(path)}")
        ps = _settings_for(path, s, dims)
        file_out = os.path.join(out_dir, *gi.folders[i]) if out_dir else None
        edits, src = find_saved_edits(path) if apply_saved_edits else ([], None)
        t1 = time.perf_counter()
        try:
            r = run_file(path, ps, file_out, save_session=save_sessions, edits=edits, edits_source=src,
                         progress=_scaled(progress, i / max(n, 1), (i + 1) / max(n, 1)))
        except KeyboardInterrupt:
            raise
        except Exception as exc:                                    # run_file records errors; this is a crash
            r = RunResult(path, ps, errors={"process": f"{type(exc).__name__}: {exc}"})
        row = {**prefix, **{k: v for k, v in summary_row(r).items() if k != "file"}}
        rows.append(row)
        if pooled_tables and r.ok:
            for tab, fn in ((fib, fibre_rows), (dom, domain_rows)):
                try:
                    tab += [{**prefix, **x} for x in fn(r)]
                except Exception as exc:                            # a table must not stop the batch
                    log.append(f"    table not made for {gi.samples[i]}: {type(exc).__name__}: {exc}")
        checks += [{**prefix, **x} for x in data_check_rows(r.data_check, row["status"])]
        results.append(r.light())
        msg = "ok" if r.ok else "; ".join(f"{k}: {v}" for k, v in r.errors.items())
        log.append(f"[{i + 1}/{n}] {gi.samples[i]}: {msg} ({time.perf_counter() - t1:.1f} s)")
        for w in r.edit_warnings:
            log.append(f"    {w}")
        del r
        if on_result:
            on_result(i, n, results[-1])
    report(progress, 1.0, "Writing the summary")
    sheets: List[Tuple[str, str, Sequence[Mapping[str, Any]]]] = [("Summary", "OTM_batch_summary", rows)]
    ok_rows = [r for r in rows if r.get("status") == "ok"]
    K = len(gi.columns)
    if group_sheets and K >= 1 and ok_rows:
        sets = [(gi.columns, "", " x ".join(gi.names), "_".join(gi.columns))]
        if K >= 2:
            sets.append((gi.columns[1:], gi.columns[0], f"{' x '.join(gi.names[1:])} (across {gi.names[0]})",
                         f"{'_'.join(gi.columns[1:])}_across_{gi.columns[0]}"))
            sets.append(([gi.columns[0]], "", gi.names[0], gi.columns[0]))
        for keys, unit, title, stem in sets:
            sheets.append((f"By {title}", f"OTM_batch_by_{stem}",
                           group_stats(ok_rows, keys, unit, f"n_{unit}" if unit else "n_units")))
    if pooled_tables:
        sheets += [("Fibres", "OTM_batch_fibres", fib), ("Capillary domains", "OTM_batch_domains", dom)]
    sheets.append(("Data checks", "OTM_batch_data_checks", checks))
    written = write_workbook(sheets, out)
    written.append(s.save(os.path.join(out, SETTINGS_FILE)))
    dt = time.perf_counter() - t0
    n_ok = len(ok_rows)
    log.append(f"Done in {dt:.0f} s: {n_ok} ok, {sum(1 for r in rows if r['status'] not in ('ok', 'not run'))} "
               f"with errors, {sum(1 for r in rows if r['status'] == 'not run')} not run.")
    lp = os.path.join(out, LOG_FILE)
    with open(lp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(log) + "\n")
    written.append(lp)
    return ExperimentResult(files, gi, results, rows, written, out, dt, stop)


def _scaled(progress, a, b):
    from .progress import scaled

    return scaled(progress, a, b)
