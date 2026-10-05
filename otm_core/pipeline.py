"""The whole OTM analysis without a window: one call per data file, or a batch.

Steps (the desktop wizard's, in order)::

    load -> retouch -> dimensions + ROI -> indices
                                       \\-> mesh -> solve (PO2) -> flux lines
    -> exports (INDICES/ and PO2/ next to the data file, or under --out)

:class:`RunSettings` holds every choice the wizard asks for; it is saved as JSON
next to the batch summary, so a run can be repeated exactly
(``python -m otm_core run --settings OTM_batch_settings.json ...``).

:func:`run_file` never raises for a bad data file or a failed step: the error is
recorded in :class:`RunResult` and the steps that do not depend on the failed one
still run (e.g. the indices when meshing fails). :func:`run_batch` runs many files
(optionally in parallel processes) and writes one summary table (``.xlsx`` and
``.csv``) with a row per file.

No GUI code here: the desktop app, the command line (:mod:`otm_core.__main__`)
and a future web back end can all call it.
"""

from __future__ import annotations

import csv
import dataclasses
import glob
import json
import math
import os
import time
import traceback
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np

from .models import ROI, Geometry, LengthScale, RetouchSettings
from .progress import ProgressCallback, report, scaled

__all__ = ["RunSettings", "RunResult", "STEPS", "run_file", "run_batch", "find_data_files", "read_dimensions_table",
           "summary_row", "write_summary", "SUMMARY_FILE", "SETTINGS_FILE"]

STEPS = ("indices", "po2", "flux")
"""Analysis steps a run can include; "flux" needs "po2"."""
EXERCISE_LEVELS = ("resting", "low", "moderate", "high")
SUMMARY_FILE = "OTM_batch_summary"
SETTINGS_FILE = "OTM_batch_settings.json"


# ==========================================================================
# settings
# ==========================================================================


def _mesh_settings_default():
    from .fem import MeshSettings

    return MeshSettings()


def _flux_settings_default():
    from .flux import FluxSettings

    return FluxSettings()


@dataclass
class RunSettings:
    """Everything the wizard asks for. Sizes in µm; ROI as (x_min, x_max, y_min, y_max)."""

    tissue: str = "skeletal"                       # "skeletal" | "cardiac"
    use_fibre_types: bool = True
    width_um: Optional[float] = None               # give the width *or* the height
    height_um: Optional[float] = None
    roi_um: Optional[Tuple[float, float, float, float]] = None   # None: 20-80 % of each side
    retouch: RetouchSettings = field(default_factory=RetouchSettings)
    parameters: Optional[Dict[str, float]] = None  # TransportParameters fields; None: tissue defaults
    parameter_file: Optional[str] = None           # legacy .dat (overrides ``parameters``)
    exercise: str = "resting"
    non_uniform: bool = True
    differential_extraction: float = 1.0
    mesh: Any = field(default_factory=_mesh_settings_default)    # otm_core.fem.MeshSettings
    flux: Any = field(default_factory=_flux_settings_default)    # otm_core.flux.FluxSettings
    steps: Tuple[str, ...] = STEPS
    index_seed: Optional[int] = 20221              # reproducible random NN pairings
    export: bool = True
    figures: bool = True
    dpi: Optional[int] = None                      # None: 600 (publication) / 300 (legacy)
    colormap: str = "turbo"
    figure_layout: str = "publication"             # "publication" (journal-ready) | "legacy" (MATLAB look)
    figure_width: str = "double"                   # "single" (85 mm) | "onehalf" (114 mm) | "double" (175 mm)
    figure_formats: Tuple[str, ...] = ("png", "pdf")
    po2_range_mmHg: Optional[Tuple[float, float]] = None

    def __post_init__(self):
        if self.tissue not in ("skeletal", "cardiac"):
            raise ValueError("tissue must be 'skeletal' or 'cardiac'")
        if self.exercise.lower() not in EXERCISE_LEVELS:
            raise ValueError(f"exercise must be one of {EXERCISE_LEVELS}")
        self.exercise = self.exercise.lower()
        self.steps = tuple(self.steps)
        bad = set(self.steps) - set(STEPS)
        if bad:
            raise ValueError(f"unknown steps {sorted(bad)}; choose from {STEPS}")
        if "flux" in self.steps and "po2" not in self.steps:
            self.steps = tuple(s for s in STEPS if s in self.steps or s == "po2")
        if self.width_um is not None and self.height_um is not None:
            raise ValueError("give the tissue width or its height, not both (the other follows the image)")
        for v in (self.width_um, self.height_um):
            if v is not None and not float(v) > 0:
                raise ValueError("the tissue size must be positive")
        self.figure_formats = tuple(self.figure_formats)
        self.figure_style()                                      # validates layout, width and formats
        if self.roi_um is not None:
            r = tuple(float(v) for v in self.roi_um)
            if len(r) != 4 or not (r[0] < r[1] and r[2] < r[3]):
                raise ValueError("roi_um must be (x_min, x_max, y_min, y_max) with min < max")
            self.roi_um = r

    # ---- derived choices -------------------------------------------------------------------------
    def transport_parameters(self):
        from .fem import TransportParameters

        if self.parameter_file:
            return TransportParameters.from_dat(self.parameter_file)
        if self.parameters:
            return dataclasses.replace(TransportParameters.defaults(self.tissue), **self.parameters)
        return TransportParameters.defaults(self.tissue)

    def figure_style(self):
        from .figures import FigureStyle

        if self.figure_layout == "legacy":
            return FigureStyle.legacy(formats=self.figure_formats, dpi=self.dpi or 300)
        return FigureStyle(layout=self.figure_layout, width=self.figure_width, formats=self.figure_formats,
                           dpi=self.dpi or 600)

    def switches(self):
        from .fem import ModelSwitches

        if self.tissue == "cardiac":
            return ModelSwitches.cardiac(self.exercise)
        return ModelSwitches.skeletal(self.exercise, self.non_uniform, self.differential_extraction)

    def with_changes(self, **changes: Any) -> "RunSettings":
        return dataclasses.replace(self, **changes)

    # ---- JSON ---------------------------------------------------------------------------------------
    def to_dict(self) -> Dict[str, Any]:
        d = {}
        for f in dataclasses.fields(self):
            v = getattr(self, f.name)
            if dataclasses.is_dataclass(v):
                v = dataclasses.asdict(v)
            elif isinstance(v, tuple):
                v = list(v)
            d[f.name] = v
        return d

    @classmethod
    def from_dict(cls, d: Mapping[str, Any]) -> "RunSettings":
        from .fem import MeshSettings
        from .flux import FluxSettings

        known = {f.name for f in dataclasses.fields(cls)}
        unknown = set(d) - known
        if unknown:
            raise ValueError(f"unknown settings: {sorted(unknown)}")
        kw = dict(d)
        for name, typ in (("retouch", RetouchSettings), ("mesh", MeshSettings), ("flux", FluxSettings)):
            if isinstance(kw.get(name), Mapping):
                kw[name] = _dataclass_from(typ, kw[name])
        for name in ("roi_um", "po2_range_mmHg", "steps", "figure_formats"):
            if kw.get(name) is not None:
                kw[name] = tuple(kw[name])
        return cls(**kw)

    def save(self, path: str) -> str:
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(self.to_dict(), fh, indent=2)
        return path

    @classmethod
    def load(cls, path: str) -> "RunSettings":
        with open(path, "r", encoding="utf-8") as fh:
            return cls.from_dict(json.load(fh))


def _dataclass_from(typ: Any, d: Mapping[str, Any]) -> Any:
    names = {f.name for f in dataclasses.fields(typ)}
    unknown = set(d) - names
    if unknown:
        raise ValueError(f"unknown {typ.__name__} settings: {sorted(unknown)}")
    return typ(**{k: (tuple(v) if isinstance(v, list) else v) for k, v in d.items()})


# ==========================================================================
# one file
# ==========================================================================


@dataclass
class RunResult:
    """Outcome of :func:`run_file`. ``errors`` maps a step to its message; a step
    missing from ``timings`` did not run (not requested, or its input failed)."""

    path: str
    settings: RunSettings
    base_name: str = ""
    raw: Any = None                        # otm_core.io.DtectData
    ingest: Any = None
    geometry: Optional[Geometry] = None
    morphometrics: Any = None
    mesh: Any = None                       # otm_core.fem.TissueMesh
    model: Any = None                      # otm_core.fem.ModelGeometry
    solution: Any = None
    flux_lines: Any = None
    files: List[str] = field(default_factory=list)
    edits: List[Any] = field(default_factory=list)         # manual corrections applied (otm_core.edits.Edit)
    edit_warnings: List[str] = field(default_factory=list)
    timings: Dict[str, float] = field(default_factory=dict)
    errors: Dict[str, str] = field(default_factory=dict)
    details: Dict[str, str] = field(default_factory=dict)      # step -> traceback

    @property
    def ok(self) -> bool:
        return not self.errors

    def summary(self) -> Dict[str, Any]:
        return summary_row(self)

    def to_session(self) -> Any:
        """:class:`otm_core.session.Session` of this run (open it in the desktop app)."""
        from .session import Session

        s = self.settings
        g = self.geometry
        return Session(raw=self.raw, source_path=self.path, tissue=s.tissue, use_fibre_types=s.use_fibre_types,
                       retouch_settings=s.retouch, length_scale=g.length_scale if g is not None else None,
                       roi=g.roi if g is not None else None, parameters=s.transport_parameters(),
                       switches=s.switches(), mesh_settings=s.mesh, flux_settings=s.flux, index_seed=s.index_seed,
                       geometry=g, has_indices=self.morphometrics is not None, mesh=self.mesh, model=self.model,
                       solution=self.solution, flux_lines=self.flux_lines,
                       view={"colormap": s.colormap, "po2_range_mmHg": s.po2_range_mmHg}, edits=list(self.edits))

    def light(self) -> "RunResult":
        """Copy without the large arrays (what a worker process sends back);
        its :meth:`summary` was computed before they were dropped."""
        return _Light(self)


class _Light(RunResult):
    """A RunResult whose summary was computed before the arrays were dropped."""

    def __init__(self, r: RunResult):
        super().__init__(path=r.path, settings=r.settings, base_name=r.base_name, files=list(r.files),
                         edits=list(r.edits), edit_warnings=list(r.edit_warnings),
                         timings=dict(r.timings), errors=dict(r.errors), details=dict(r.details))
        self._row = summary_row(r)

    def summary(self) -> Dict[str, Any]:
        return dict(self._row)


def _output_folders(path: str, out_dir: Optional[str], stem: str) -> Tuple[str, str]:
    """``INDICES``/``PO2`` next to the data file, or ``<out>/<stem>/INDICES`` ..."""
    root = os.path.join(out_dir, stem) if out_dir else os.path.dirname(os.path.abspath(path))
    return os.path.join(root, "INDICES"), os.path.join(root, "PO2")


def run_file(path: str, settings: Optional[RunSettings] = None, out_dir: Optional[str] = None,
             base_name: Optional[str] = None, progress: Optional[ProgressCallback] = None,
             save_session: bool = False, edits: Optional[Sequence[Any]] = None,
             edits_source: Optional[Dict[str, Any]] = None) -> RunResult:
    """Run the analysis on one Dtect export. Never raises for analysis errors.

    ``save_session``: also write ``<name>.otm`` (with the mesh, PO2 and flux lines)
    in the output folder, to open the result in the desktop app.
    ``edits``: manual corrections (:class:`otm_core.edits.Edit`) to apply after
    loading; ``edits_source`` = ``session.json`` they come from, whose ``raw``
    block must describe the same data (else the file fails with an error)."""
    from . import get_morphometric_data, ingest_dtect, retouch_fibers
    from .io import load_dtect_mat

    s = settings or RunSettings()
    stem = os.path.splitext(os.path.basename(path))[0]
    res = RunResult(path=os.path.abspath(path), settings=s, base_name=base_name or stem, edits=list(edits or []))
    plan = [("load", 0.04), ("geometry", 0.08)]
    if "indices" in s.steps:
        plan.append(("indices", 0.12))
    if "po2" in s.steps:
        plan += [("mesh", 0.35), ("po2", 0.15)]
    if "flux" in s.steps:
        plan.append(("flux", 0.12))
    if s.export:
        plan.append(("export", 0.14))
    if save_session:
        plan.append(("session", 0.04))
    total = sum(w for _, w in plan)
    start = {}
    acc = 0.0
    for name, w in plan:
        start[name] = (acc / total, (acc + w) / total)
        acc += w

    def step(name: str, fn: Callable[[Optional[ProgressCallback]], Any]) -> Any:
        a, b = start[name]
        report(progress, a, name)
        t0 = time.perf_counter()
        try:
            return fn(scaled(progress, a, b))
        except KeyboardInterrupt:
            raise
        except Exception as exc:                                  # recorded, the batch goes on
            res.errors[name] = f"{type(exc).__name__}: {exc}"
            res.details[name] = traceback.format_exc()
            return None
        finally:
            res.timings[name] = time.perf_counter() - t0

    # ---- load + ingest
    def load(_p):
        data = load_dtect_mat(path)
        res.raw = data
        skeletal = s.tissue == "skeletal"
        fibres = data.fibers_px if skeletal else []
        types = data.fiber_types if (skeletal and s.use_fibre_types) else None
        ing = ingest_dtect(data.capillaries_px, fibres, data.image_size, types)
        if res.edits:
            from .edits import apply_edits

            if edits_source is not None:
                rb = edits_source.get("raw", {})
                same = (list(rb.get("image_size", [])) == list(map(int, data.image_size))
                        and rb.get("n_capillaries") in (None, len(data.capillaries_px))
                        and rb.get("n_fibres") in (None, len(data.fibers_px)))
                if not same:
                    raise ValueError("the saved corrections were made on different data "
                                     f"({edits_source.get('_path', 'session')})")
            er = apply_edits(ing, res.edits)
            res.edit_warnings = er.warnings
            ing = er.ingest
        return ing

    res.ingest = step("load", load)
    if res.ingest is None:
        return _finish(res, progress)

    # ---- retouch, dimensions, ROI
    def geometry(_p):
        ing = res.ingest
        if s.width_um is None and s.height_um is None:
            raise ValueError("tissue size missing: give the width or height in µm (--width/--height or --dims)")
        ls = (LengthScale.from_width(float(s.width_um), ing.frame) if s.width_um is not None
              else LengthScale.from_height(float(s.height_um), ing.frame))
        roi = ROI(*s.roi_um) if s.roi_um is not None else ROI.default(ls)
        if not (0 <= roi.x_min_um < roi.x_max_um <= ls.x_um + 1e-9 and 0 <= roi.y_min_um < roi.y_max_um
                <= ls.y_um + 1e-9):
            raise ValueError(f"the ROI must lie inside the tissue (0-{ls.x_um:g} x 0-{ls.y_um:g} µm)")
        rt = retouch_fibers(s.retouch, ing.fibers_pxc, ing.capillaries_pxc, ing.frame)
        types = ing.fiber_types if len(rt.rescaled) else np.zeros(0, dtype=int)
        return Geometry(ing.frame, rt.capillaries_ndim, list(rt.rescaled), types, ls, roi)

    res.geometry = step("geometry", geometry)
    if res.geometry is None:
        return _finish(res, progress)
    g = res.geometry

    # ---- indices
    if "indices" in s.steps:
        res.morphometrics = step("indices", lambda p: get_morphometric_data(
            g, rng=np.random.default_rng(s.index_seed), emit_warnings=False))

    # ---- mesh, PO2, flux
    if "po2" in s.steps:
        def mesh(p):
            from .fem import mesh_tissue

            params = s.transport_parameters().derive(len(g.capillaries), g.length_scale)
            return params, mesh_tissue(g, params.Rcap, s.mesh, progress=p)

        out = step("mesh", mesh)
        if out is not None:
            params, (res.mesh, res.model) = out

            def po2(p):
                from .fem import solve_po2

                sol = solve_po2(res.mesh, params, s.switches(), tissue=s.tissue, progress=p)
                if not sol.converged:
                    raise RuntimeError(f"Newton did not converge in {sol.iterations} iterations")
                return sol

            res.solution = step("po2", po2)

    if "flux" in s.steps and res.solution is not None:
        def flux(p):
            from .flux import flux_lines_from_solution, roi_capillary_index
            from .morphometry import capillary_domain_statistics, voronoi_cells

            caps = np.asarray(g.capillaries, float)
            if res.morphometrics is not None:
                roi_idx = res.morphometrics.domains.roi_index
            else:
                roi_idx = capillary_domain_statistics(caps, g.length_scale, g.roi_ndim(), voronoi_cells(caps)).roi_index
            roi = roi_capillary_index(caps, roi_idx, res.mesh.capillaries)
            if len(roi) == 0:
                raise RuntimeError("no capillaries in the ROI")
            return flux_lines_from_solution(res.solution, roi, s.flux, progress=p)

        res.flux_lines = step("flux", flux)

    # ---- exports
    if s.export:
        def export(p):
            from . import export as ex

            ind_dir, po2_dir = _output_folders(path, out_dir, stem)
            files: List[str] = []
            fs = s.figure_style()
            if res.morphometrics is not None:
                files += ex.export_indices(res.morphometrics, g, ind_dir, res.base_name, s.figures, s.dpi,
                                           progress=scaled(p, 0.0, 0.45), style=fs)
            if res.solution is not None:
                files += ex.export_po2(res.solution, g, po2_dir, res.base_name, s.tissue,
                                       getattr(res.model, "fiber_index", None), s.figures, s.colormap,
                                       s.po2_range_mmHg, s.dpi, progress=scaled(p, 0.45, 0.9), style=fs)
            if res.flux_lines is not None:
                vor = res.morphometrics.domains.voronoi if res.morphometrics is not None else None
                files += ex.export_flux(res.flux_lines, g, res.mesh, po2_dir, res.base_name, vor,
                                        getattr(res.model, "capillary_index", None), s.figures, s.dpi,
                                        progress=scaled(p, 0.9, 1.0), style=fs)
            return files

        res.files = step("export", export) or []
    if save_session:
        def session(_p):
            from .session import save_session as _save

            from .session import SessionError, read_session_info

            root = os.path.join(out_dir, stem) if out_dir else os.path.dirname(os.path.abspath(path))
            os.makedirs(root, exist_ok=True)
            target = os.path.join(root, res.base_name + ".otm")
            if os.path.isfile(target):              # never overwrite hand-made corrections
                try:
                    old = read_session_info(target).get("edits") or []
                except SessionError:
                    old = []
                if old and old != [e.to_dict() for e in res.edits]:
                    target = os.path.join(root, res.base_name + "_batch.otm")
            return _save(target, res.to_session())

        f = step("session", session)
        if f:
            res.files.append(f)
    return _finish(res, progress)


def _finish(res: RunResult, progress: Optional[ProgressCallback]) -> RunResult:
    report(progress, 1.0, "done" if res.ok else "finished with errors")
    return res


# ==========================================================================
# summary table
# ==========================================================================


def _m(stats: Any) -> float:
    return float(stats.mean) if stats is not None and np.isfinite(stats.mean) else float("nan")


def summary_row(res: RunResult) -> Dict[str, Any]:
    """One row of the batch table (flat; NaN where a step did not run)."""
    nan = float("nan")
    g, m, sol, fl = res.geometry, res.morphometrics, res.solution, res.flux_lines
    row: Dict[str, Any] = {"file": os.path.basename(res.path), "folder": os.path.dirname(res.path),
                           "status": "ok" if res.ok else "error",
                           "errors": "; ".join(f"{k}: {v}" for k, v in res.errors.items())}
    s = res.settings
    row.update({"tissue": s.tissue, "exercise": s.exercise, "corrections": len(res.edits),
                "corrections_skipped": len(res.edit_warnings)})
    ls = g.length_scale if g is not None else None
    row.update({"width_um": ls.x_um if ls else nan, "height_um": ls.y_um if ls else nan,
                "roi_x_min_um": g.roi.x_min_um if g is not None else nan,
                "roi_x_max_um": g.roi.x_max_um if g is not None else nan,
                "roi_y_min_um": g.roi.y_min_um if g is not None else nan,
                "roi_y_max_um": g.roi.y_max_um if g is not None else nan,
                "n_capillaries": len(g.capillaries) if g is not None else nan,
                "n_fibres": len(g.fibers) if g is not None else nan})
    d = m.domains if m is not None else None
    nn = m.nearest_neighbours if m is not None else None
    sup = m.supply if m is not None else None
    row.update({
        "capillary_density_per_mm2": d.capillary_density_per_mm2 if d else nan,
        "roi_capillary_density_per_mm2": d.roi_capillary_density_per_mm2 if d else nan,
        "domains_in_roi": d.n_domains_in_roi if d else nan,
        "domain_area_mean_um2": d.area_stats.mean if d else nan,
        "domain_area_sd_um2": d.area_stats.std if d else nan,
        "domain_area_logsd": _logsd(d.roi_capillary_domain_areas_um2) if d else nan,
        "domain_eq_diameter_mean_um": d.diameter_stats.mean if d else nan,
        "nn_mean_um": nn.summary["MeanOfMeans"] if nn else nan,
        "nn_min_mean_um": nn.summary["MeanOfMin"] if nn else nan,
        "fibres_in_roi": sup.number_of_fibers if sup else nan,
        "capillary_to_fibre_ratio": sup.capillary_to_fiber_ratio if sup else nan,
        "fibre_area_mean_um2": float(np.mean(sup.fiber_area_um2)) if sup is not None and len(sup.fiber_area_um2)
        else nan,
    })
    for name in ("LCFR", "LCD", "DFR", "FDR", "Dmax", "CC", "CFi", "FPi", "CFPE", "SF"):
        st = getattr(sup, name) if sup is not None else None
        row[f"{name}_mean"] = _m(st)
        row[f"{name}_sd"] = float(st.std) if st is not None else nan
    row.update({"mesh_nodes": res.mesh.n_nodes if res.mesh is not None else nan,
                "mesh_triangles": res.mesh.n_triangles if res.mesh is not None else nan})
    if sol is not None:
        from .fem import mo2_statistics, po2_statistics

        ps, ms = po2_statistics(sol), mo2_statistics(sol)
        for key, lab in (("Tissue", "tissue"), ("Interstitia", "is"), ("AllFibers", "fibres"),
                         ("FiberI", "type_I"), ("FiberIIa", "type_IIa"), ("FiberIIb", "type_IIb")):
            v = ps.get(key, {})
            row[f"po2_{lab}_mean_mmHg"] = v.get("mean_mmHg", nan)
            row[f"po2_{lab}_sd_mmHg"] = v.get("std_mmHg", nan)
            row[f"po2_{lab}_hypoxic_pct"] = v.get("hypoxia_pct", nan)
        t = ms.get("Tissue", {})
        row.update({"mo2_tissue_mean": t.get("mean_mo2", nan), "vo2max_tissue": t.get("vo2max", nan),
                    "newton_iterations": sol.iterations})
        row["po2_min_mmHg"] = sol.summary()["po2_min_mmHg"]
    if fl is not None:
        um = ls.y_um / 2.0 if ls else nan
        sm = fl.summary()
        L = fl.lengths()
        row.update({"flux_seed_capillaries": sm["n_seed_capillaries"],
                    "flux_capillaries_with_lines": sm["n_capillaries_with_lines"],
                    "flux_moving_lines": sm["n_moving_lines"],
                    "flux_median_length_um": float(np.median(L)) * um if L.size else nan,
                    "flux_max_length_um": float(fl.max_length) * um})
    for k, v in res.timings.items():
        row[f"seconds_{k}"] = round(v, 2)
    row["files_written"] = len(res.files)
    return row


def _logsd(v: np.ndarray) -> float:
    v = np.asarray(v, float)
    return float(np.std(np.log10(v), ddof=1)) if v.size > 1 else float("nan")


def write_summary(rows: Sequence[Mapping[str, Any]], stem: str) -> List[str]:
    """``<stem>.csv`` and ``<stem>.xlsx`` (when openpyxl is installed); columns in
    the order they first appear, so files missing a step get blanks."""
    cols: List[str] = []
    for r in rows:
        cols += [c for c in r if c not in cols]
    os.makedirs(os.path.dirname(os.path.abspath(stem)) or ".", exist_ok=True)
    out = []
    with open(stem + ".csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for r in rows:
            w.writerow([_csv(r.get(c)) for c in cols])
    out.append(stem + ".csv")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:                               # pragma: no cover
        return out
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_xl(r.get(c)) for c in cols])
    ws.freeze_panes = "C2"
    for k, name in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, k).column_letter].width = min(40, max(10, len(name) + 2))
    wb.save(stem + ".xlsx")
    out.append(stem + ".xlsx")
    return out


def _csv(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, float):
        return "" if not math.isfinite(v) else repr(v)
    return v


def _xl(v: Any) -> Any:
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return float(v) if math.isfinite(v) else None
    return v


# ==========================================================================
# many files
# ==========================================================================


def find_data_files(patterns: Iterable[str]) -> List[str]:
    """Expand files, folders (their ``*.mat``) and wildcards (``**`` recursive);
    sorted, without duplicates. Wildcards are expanded here because the Windows
    shell does not do it."""
    out: List[str] = []
    for p in patterns:
        if os.path.isdir(p):
            hits = sorted(glob.glob(os.path.join(p, "*.mat")))
        elif any(ch in p for ch in "*?["):
            hits = sorted(glob.glob(p, recursive=True))
        else:
            hits = [p]
        for h in hits:
            a = os.path.abspath(h)
            if a not in out:
                out.append(a)
    return out


def read_dimensions_table(path: str) -> Dict[str, Dict[str, Any]]:
    """Per-file sizes from a CSV with a ``file`` column (name, with or without
    ``.mat``) and ``width_um`` or ``height_um``; optional ``roi_x_min_um``,
    ``roi_x_max_um``, ``roi_y_min_um``, ``roi_y_max_um``. Returns
    ``{file stem: {settings changes}}``."""
    out: Dict[str, Dict[str, Any]] = {}
    with open(path, newline="", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            r = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
            name = r.get("file", "")
            if not name:
                continue
            ch: Dict[str, Any] = {}
            if r.get("width_um"):
                ch["width_um"], ch["height_um"] = float(r["width_um"]), None
            elif r.get("height_um"):
                ch["height_um"], ch["width_um"] = float(r["height_um"]), None
            roi = [r.get(k) for k in ("roi_x_min_um", "roi_x_max_um", "roi_y_min_um", "roi_y_max_um")]
            if all(roi):
                ch["roi_um"] = tuple(float(v) for v in roi)
            out[os.path.splitext(os.path.basename(name))[0].lower()] = ch
    return out


def _settings_for(path: str, settings: RunSettings, dims: Optional[Mapping[str, Mapping[str, Any]]]) -> RunSettings:
    if not dims:
        return settings
    ch = dims.get(os.path.splitext(os.path.basename(path))[0].lower())
    return settings.with_changes(**ch) if ch else settings


def find_saved_edits(path: str) -> Tuple[List[Any], Optional[Dict[str, Any]]]:
    """Manual corrections saved by the desktop app for ``path``: the session
    ``<name>.otm`` next to the data file. ([], None) when there is none."""
    from .edits import edits_from_list
    from .session import SessionError, read_session_info

    cand = os.path.splitext(path)[0] + ".otm"
    if not os.path.isfile(cand):
        return [], None
    try:
        meta = read_session_info(cand)
    except SessionError:
        return [], None
    meta["_path"] = os.path.basename(cand)
    return edits_from_list(meta.get("edits") or []), meta


def _worker(args: Tuple[str, RunSettings, Optional[str], bool, bool]) -> RunResult:
    path, s, out_dir, sessions, use_edits = args
    edits, src = find_saved_edits(path) if use_edits else ([], None)
    return run_file(path, s, out_dir, save_session=sessions, edits=edits, edits_source=src).light()


def run_batch(paths: Sequence[str], settings: Optional[RunSettings] = None, out_dir: Optional[str] = None,
              summary_stem: Optional[str] = None, dims: Optional[Mapping[str, Mapping[str, Any]]] = None,
              jobs: int = 1, save_sessions: bool = False, apply_saved_edits: bool = False, on_result: Optional[Callable[[int, int, RunResult], None]] = None,
              progress: Optional[ProgressCallback] = None) -> Tuple[List[RunResult], List[str]]:
    """Run every file; returns (results in input order, summary files written).

    ``jobs > 1`` runs files in separate processes (each holds one mesh at a time;
    results come back without the large arrays). ``on_result(i, n, result)`` is
    called as each file finishes. The settings are saved next to the summary.
    ``apply_saved_edits``: apply the manual corrections saved by the desktop app
    in ``<name>.otm`` next to each data file (:func:`find_saved_edits`).
    """
    s = settings or RunSettings()
    paths = list(paths)
    n = len(paths)
    per_file = [_settings_for(p, s, dims) for p in paths]
    results: List[Optional[RunResult]] = [None] * n
    if jobs > 1 and n > 1:
        from concurrent.futures import ProcessPoolExecutor, as_completed

        with ProcessPoolExecutor(max_workers=min(jobs, n)) as pool:
            futs = {pool.submit(_worker, (p, ps, out_dir, save_sessions, apply_saved_edits)): i for i, (p, ps) in enumerate(zip(paths, per_file))}
            done = 0
            for f in as_completed(futs):
                i = futs[f]
                try:
                    results[i] = f.result()
                except Exception as exc:                         # the worker process itself died
                    results[i] = RunResult(paths[i], per_file[i], errors={"process": f"{type(exc).__name__}: {exc}"})
                done += 1
                report(progress, done / n, os.path.basename(paths[i]))
                if on_result:
                    on_result(i, n, results[i])
    else:
        for i, (p, ps) in enumerate(zip(paths, per_file)):
            edits, src = find_saved_edits(p) if apply_saved_edits else ([], None)
            r = run_file(p, ps, out_dir, progress=scaled(progress, i / n, (i + 1) / n), save_session=save_sessions,
                         edits=edits, edits_source=src)
            results[i] = r.light()
            del r                                               # free the mesh before the next file
            if on_result:
                on_result(i, n, results[i])
    written: List[str] = []
    if summary_stem:
        written = write_summary([r.summary() for r in results], summary_stem)
        written.append(s.save(os.path.join(os.path.dirname(os.path.abspath(summary_stem)), SETTINGS_FILE)))
    return [r for r in results if r is not None], written
